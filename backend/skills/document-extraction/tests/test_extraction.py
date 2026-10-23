import importlib.util
import json
from pathlib import Path
import pytest

SPEC = importlib.util.spec_from_file_location("document_extraction_skill", Path(__file__).parents[1] / "scripts" / "skill.py")
skill = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(skill)


def test_pdf_embedded_bill_of_lading_is_confirmed_by_the_layout_line():
    value = 'CGT260801188'
    embedded = 'HOUSE B/L NO. : ' + value
    document = dict(source_id='pdf', document_id='doc-1', bill_of_lading=value,
                    evidence=[dict(source_id='pdf', anchor='page:1/embedded',
                                   quote=embedded, kind='pdf_embedded_text',
                                   field='bill_of_lading')])
    source = dict(source_id='pdf', format='pdf', status='ready', blocks=[
        dict(anchor='page:1/embedded', kind='pdf_embedded_text', text=embedded),
        dict(anchor='page:1/layout', kind='pdf_layout_text',
             text='HOUSE B/L NO. : ' + value + '   POD / ETA')])
    result = skill.validate_extraction(dict(result=dict(documents=[document]), sources=[source]))
    assert not any(i['code'] == 'PDF_IDENTITY_LAYOUT_REQUIRED' for i in result['issues'])
    anchors = {ref['anchor'] for ref in result['documents'][0]['evidence']
               if ref.get('field') == 'bill_of_lading'}
    assert 'page:1/layout' in anchors


def test_document_identity_evidence_is_repaired_from_the_same_source():
    value = 'TIIU8975975'
    document = dict(source_id='mail', document_id='doc-1', container_no=value,
                    evidence=[dict(source_id='mail', anchor='mail/headers',
                                   quote='From: luka@pltplt.com', field='container_no')])
    source = dict(source_id='mail', format='eml', status='ready', blocks=[
        dict(anchor='mail/headers', kind='mail_headers',
             text='Subject: PRE-ALERT CNTR NO.: ' + value + '\nFrom: luka@pltplt.com'),
        dict(anchor='mail/body', kind='mail_body', text='SHA pay TAR：H/C $50')])
    result = skill.validate_extraction(dict(result=dict(documents=[document]), sources=[source]))
    refs = [ref for ref in result['documents'][0]['evidence']
            if ref.get('field') == 'container_no']
    assert refs and value in refs[0]['quote']
    assert any(i['code'] == 'DOCUMENT_EVIDENCE_REPAIRED' for i in result['issues'])


def test_duplicate_issues_are_collapsed():
    duplicate = dict(code='DUPLICATE', field='x', message='same')
    result = skill.validate_extraction(dict(result=dict(issues=[duplicate, dict(duplicate)])))
    assert sum(i.get('code') == 'DUPLICATE' for i in result['issues']) == 1


def test_platform_footer_vendor_is_never_the_issuer():
    vendor = '湃乐多（上海）航运科技有限公司 / Pilot Aiship Tech co., Ltd.'
    result = skill.validate_extraction(dict(result=dict(
        charges=[dict(source_id='pdf', issuer=vendor, charge_item='HANDLING CHARGE',
                      amount=65, currency='USD', evidence=[dict(source_id='pdf',
                      anchor='page:3/layout', quote='Powered by ' + vendor, field='issuer')])],
        documents=[dict(source_id='pdf', document_id='doc-1', issuer=vendor)]),
        sources=[dict(source_id='pdf', format='pdf', status='ready', blocks=[])]))
    assert result['charges'][0]['issuer'] is None
    assert result['documents'][0]['issuer'] is None
    assert sum(i['code'] == 'ISSUER_PLATFORM_VENDOR' for i in result['issues']) == 2


def test_printed_issuer_is_kept():
    result = skill.validate_extraction(dict(result=dict(charges=[dict(
        source_id='pdf', issuer='NEPTUNE SHIPPING LIMITED', charge_item='HANDLING CHARGE',
        amount=65, currency='USD', evidence=[dict(source_id='pdf', anchor='page:1/layout',
        quote='BENEFICIARY : NEPTUNE SHIPPING LIMITED', field='issuer')])]),
        sources=[dict(source_id='pdf', format='pdf', status='ready', blocks=[])]))
    assert result['charges'][0]['issuer'] == 'NEPTUNE SHIPPING LIMITED'
    assert not any(i['code'] == 'ISSUER_PLATFORM_VENDOR' for i in result['issues'])


def test_read_materials_accepts_zero_start_and_lists_valid_sources():
    source = dict(source_id='s', name='x', status='ready', format='text',
                  blocks=[dict(anchor='body', text='line1\nline2', authoritative=True)])
    result = skill.read_materials(dict(sources=[source],
        windows=[dict(source_id='s', start_line=0, end_line=2)]))
    assert result['content'][0]['start_line'] == 1
    assert 'line1' in result['content'][0]['text'] and 'line2' in result['content'][0]['text']
    with pytest.raises(ValueError, match='可用来源'):
        skill.read_materials(dict(sources=[source],
            windows=[dict(source_id='other', start_line=1)]))


def test_one_exact_quote_can_support_several_fields_without_repeating_it():
    result = skill.validate_extraction(dict(result=dict(charges=[dict(
        source_id='mail', charge_item='H/C', amount=50,
        evidence=[dict(source_id='mail', anchor='mail/body', quote='SHA pay TAR: H/C $50',
                       fields=['charge_item', 'amount', 'debtor', 'creditor'])])],
        coverage=[dict(source_id='mail', anchor='mail/body', chunk_id='a',
                       status='extracted', reason='费用已提取')]),
        chunks=[dict(chunk_id='a', source_id='mail', anchor='mail/body')]))
    assert result['valid']
    refs = result['charges'][0]['evidence']
    assert {r['field'] for r in refs} == {'charge_item', 'amount', 'debtor', 'creditor'}
    assert all(r['quote'] == 'SHA pay TAR: H/C $50' for r in refs)


def test_markdown_window_is_exact_and_eof_does_not_claim_previous_lines_read(tmp_path):
    import importlib.util
    spec = importlib.util.spec_from_file_location('md_adapter_test',
        Path(__file__).parents[2]/'document-intake/scripts/pdf_ocr_md.py')
    md = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(md)
    original = 'Page one header\r\nFreight EUR 1.234,56  \r\nLast line\r\n'
    source = dict(source_id='src-test', name='bill.txt', status='ready', format='text',
        blocks=[dict(anchor='body', text=original, authoritative=True)])
    source['markdown'] = md.build_markdown(source, tmp_path)
    line = next(s['line'] for s in source['markdown']['spans'] if 'Freight' in s['quote'])
    result = skill.read_markdown_evidence(dict(sources=[source], source_id='src-test', start_line=line, limit=1))
    assert result['text'] == 'Freight EUR 1.234,56'
    assert result['evidence'][0]['quote_start'] == original.index('Freight')
    assert result['evidence'][0]['quote'] == original[result['evidence'][0]['quote_start']:result['evidence'][0]['quote_end']]
    assert not result['eof'] and result['next_line'] == line + 1
    last = source['markdown']['line_count']
    tail = skill.read_markdown_evidence(dict(sources=[source], source_id='src-test', start_line=last, limit=1))
    assert tail['eof'] and tail['start_line'] > 1
    with pytest.raises(ValueError, match='超过文件末尾'):
        skill.read_markdown_evidence(dict(sources=[source], source_id='src-test', start_line=last + 1))
    source['blocks'][0]['text'] = original.replace('Freight', 'CHANGED')
    with pytest.raises(ValueError, match='映射与来源不一致'):
        skill.read_markdown_evidence(dict(sources=[source], source_id='src-test', start_line=line, limit=1))


def test_extraction_has_no_multimodal_page_or_image_reader():
    assert not hasattr(skill, 'read_image')
    assert not hasattr(skill, 'read_pdf_page')


def test_agent_can_request_pdf_text_without_repeating_ocr_or_claiming_failed_ocr_ready():
    source = dict(source_id='pdf', content_id='sha256:original', format='pdf', status='ready',
        blocks=[dict(anchor='page:1/native', kind='native_preview', page=1,
            authoritative=False, text='B/L NO HBL-1')],
        metadata=dict(page_count=1, ocr_ranges=[dict(pages=[1], status='ready')]))
    result = skill.supplement_pdf(dict(source=source, method='text', reason='OCR未提供提单编号'))
    block = next(b for b in result['source']['blocks'] if b['kind']=='pdf_embedded_text')
    assert block['authoritative'] and block['page'] == 1 and block['source_id'] == 'pdf'
    assert block['source_content_id'] == source['content_id']
    assert result['source']['metadata']['supplement_reads'][-1]['reason'] == 'OCR未提供提单编号'
    assert len(source['blocks']) == 1
    source['status'] = 'waiting_url'
    source['metadata']['ocr_ranges'][0]['status'] = 'waiting_url'
    result = skill.supplement_pdf(dict(source=source, method='text', reason='查看可用文字'))
    assert result['source']['status'] == 'waiting_url'
    assert not any(b.get('authoritative') for b in result['source']['blocks'])


def test_mail_transport_date_cannot_fill_missing_bill_date():
    source = dict(source_id='mail', format='eml', blocks=[
        dict(anchor='mail/headers', text='Date: 2026-10-20'),
        dict(anchor='mail/body', text='Handling charge USD 50. Invoice date: 2026-10-19')])
    charge = dict(charge_id='fee', source_id='mail', document_date='2026-10-20',
        evidence=[dict(source_id='mail', anchor='mail/headers', quote='2026-10-20', field='document_date')])
    payload = dict(sources=[source], charges=[charge], coverage=[
        dict(source_id='mail', anchor=b['anchor'], status='extracted', reason='已读') for b in source['blocks']])
    def date_errors(value):
        return [i for i in skill.validate_extraction(value)['issues'] if i.get('code') == 'BILL_DATE_EVIDENCE_REQUIRED']
    assert date_errors(payload)
    charge['human_confirmed'] = {'document_info.document_date': {'value': '2026-10-20'}}
    assert date_errors(payload), '模型自报确认不构成人工来源'
    charge['document_date'] = None
    assert not date_errors(payload), '未提供日期不转人工'
    charge['document_date'] = '2026-10-19'
    charge['evidence'] = [dict(source_id='mail', anchor='mail/body', quote='Invoice date: 2026-10-19', field='document_date')]
    assert not date_errors(payload)
    charge['document_date'], charge['evidence'] = '2026-10-20', []
    payload['records'] = [dict(charge_id='fee', document_info=dict(document_date='2026-10-20'),
        human_confirmed={'document_info.document_date': {'value': '2026-10-20'}})]
    assert not date_errors(payload), '保留宿主保存的显式人工更正'


def test_bill_date_error_points_to_candidate_lines_without_attaching_evidence():
    source = dict(source_id='notice', format='pdf', blocks=[
        dict(anchor='page:1/layout', kind='pdf_layout_text',
             text='GENERAL    NOTICE                     10/19/2026 05:58 PM\nPOD ETA: 10/26/2026')])
    charge = dict(charge_id='fee', source_id='notice', document_date='2026-10-19',
        evidence=[dict(source_id='notice', anchor='page:1/layout',
                       quote='POD ETA: 10/26/2026', field='amount')])
    payload = dict(sources=[source], charges=[charge], coverage=[
        dict(source_id='notice', anchor='page:1/layout', status='extracted', reason='已读')])

    result = skill.validate_extraction(payload)

    issue = next(issue for issue in result['issues']
                 if issue['code'] == 'BILL_DATE_EVIDENCE_REQUIRED')
    assert result['charges'][0]['document_date'] == '2026-10-19'
    assert not [ref for ref in result['charges'][0]['evidence']
                if ref.get('field') == 'document_date']
    assert issue['date_candidates'][0]['anchor'] == 'page:1/layout'
    assert '10/19/2026' in issue['date_candidates'][0]['quote']
    assert '不要仅因缺少引用就删除正确日期' in issue['message']


def test_bill_date_citation_must_print_the_claimed_date():
    source = dict(source_id='notice', format='pdf', blocks=[
        dict(anchor='page:1/layout', kind='pdf_layout_text',
             text='GENERAL NOTICE 10/19/2026\nPOD ETA: 10/26/2026')])
    charge = dict(charge_id='fee', source_id='notice', document_date='2026-10-19',
        evidence=[dict(source_id='notice', anchor='page:1/layout',
                       quote='POD ETA: 10/26/2026', field='document_date')])
    payload = dict(sources=[source], charges=[charge], coverage=[
        dict(source_id='notice', anchor='page:1/layout', status='extracted', reason='已读')])

    issues = skill.validate_extraction(payload)['issues']

    codes = [issue['code'] for issue in issues]
    assert 'BILL_DATE_EVIDENCE_MISMATCH' in codes
    assert 'BILL_DATE_EVIDENCE_REQUIRED' not in codes


def test_bill_date_without_a_source_match_still_requires_evidence():
    source = dict(source_id='notice', format='pdf', blocks=[
        dict(anchor='page:1/layout', kind='pdf_layout_text', text='POD ETA: 10/26/2026')])
    charge = dict(charge_id='fee', source_id='notice', document_date='2026-10-19',
        evidence=[dict(source_id='notice', anchor='page:1/layout',
                       quote='POD ETA: 10/26/2026', field='amount')])
    payload = dict(sources=[source], charges=[charge], coverage=[
        dict(source_id='notice', anchor='page:1/layout', status='extracted', reason='已读')])

    issues = skill.validate_extraction(payload)['issues']

    issue = next(issue for issue in issues if issue['code'] == 'BILL_DATE_EVIDENCE_REQUIRED')
    assert 'date_candidates' not in issue


def test_optional_missing_warnings_do_not_reject_an_otherwise_valid_extraction():
    payload = dict(sources=[dict(source_id='bill', text='Handling USD 50')],
        charges=[dict(source_id='bill', charge_item='Handling', amount=50, currency='USD',
                      evidence=[dict(source_id='bill', quote='Handling USD 50')])],
        coverage=[dict(source_id='bill', status='extracted', reason='已读')],
        issues=[dict(code='BILL_DATE_MISSING', severity='warning', field='document_date'),
                dict(code='NO_INVOICE_REFERENCE', severity='info', field='reference_no')])
    assert skill.validate_extraction(payload)['valid'] is True
    payload['issues'].append(dict(code='FIELD_CONFLICT', severity='error'))
    assert skill.validate_extraction(payload)['valid'] is False


def test_material_bundle_contains_every_file_without_model_authored_checklists():
    sources = [dict(source_id=s, status='ready', blocks=[dict(anchor='body', text='THC USD 50')])
               for s in ('a', 'b')]
    payload = dict(sources=sources, chunks=[dict(chunk_id='a1', source_id='a', anchor='body')],
        coverage=[dict(chunk_id='a1', source_id='a', anchor='body', status='extracted', reason='完整')])
    result = skill.read_materials(payload)
    assert {row['source_id'] for row in result['materials']} == {'a', 'b'}
    assert {row['source_id'] for row in result['content']} == {'a', 'b'}
    assert skill.validate_coverage(payload)['valid']


def test_selected_pdf_layout_uses_original_pages_hash_and_actual_header_columns(tmp_path):
    from email import policy
    from email.parser import BytesParser
    from hashlib import sha256
    from io import BytesIO
    from pypdf import PdfReader, PdfWriter

    original = next((Path(__file__).parents[3] / 'data/samples').glob('*.eml'))
    mail = BytesParser(policy=policy.default).parsebytes(original.read_bytes())
    part = next(p for p in mail.walk() if (p.get_filename() or '').endswith('MBL.pdf'))
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.add_page(PdfReader(BytesIO(part.get_payload(decode=True))).pages[0])
    path = tmp_path / 'selected-original.pdf'
    writer.write(path)
    source = dict(source_id='original', path=str(path), format='pdf', status='ready',
        content_id='sha256:' + sha256(path.read_bytes()).hexdigest(), blocks=[],
        metadata=dict(page_count=2, ocr_ranges=[dict(pages=[1, 2], status='ready')]))
    result = skill.supplement_pdf(dict(source=source, method='layout', pages=[2], reason='区分订舱与提单列'))
    assert result['added_anchors'] == ['page:2/layout']
    block = result['source']['blocks'][0]
    assert block['page'] == 2 and block['authoritative']
    assert block['source_content_id'] == source['content_id']
    lines = block['text'].splitlines()
    label = next(i for i, line in enumerate(lines) if 'BOOKING NUMBER' in line and 'B/L NO:' in line)
    assert '4055331281' in lines[label + 1] and 'HZMHK331281' in lines[label + 1]
    assert source['blocks'] == []
    reused = skill.supplement_pdf(dict(source=result['source'], method='layout', pages=[2], reason='已有版式'))
    assert reused['reused'] and reused['added_anchors'] == []
    path.write_bytes(path.read_bytes() + b'changed')
    with pytest.raises(ValueError, match='内容已变化'):
        skill.supplement_pdf(dict(source=result['source'], method='layout', pages=[2], reason='检查原件'))


@pytest.mark.parametrize('pages', [[0], [2], [True], ['1']])
def test_pdf_supplement_rejects_invalid_original_page(pages):
    source = dict(source_id='pdf', format='pdf', blocks=[], metadata=dict(page_count=1))
    with pytest.raises(ValueError, match='真实页范围'):
        skill.supplement_pdf(dict(source=source, method='text', pages=pages, reason='检查字段'))


def test_unparsed_submission_asset_cannot_disappear_from_completed_task():
    result = skill.validate_coverage(dict(asset_ids=['one', 'two'], intake=[dict(asset_id='one')],
        sources=[dict(source_id='a', text='THC USD 50')],
        coverage=[dict(source_id='a', anchor='text', status='extracted')]))
    assert any(i['code']=='ASSET_NOT_READ' for i in result['issues'])


def test_every_source_can_supply_bounded_context_and_prior_field_gaps():
    sources = [dict(source_id='a', format='pdf', blocks=[dict(anchor='p1', text='Invoice INV-1 USD 100')]),
               dict(source_id='b', format='docx', blocks=[dict(anchor='p1', text='INV-1 debtor Buyer')])]
    result = skill.prepare_extraction(dict(id='task', sources=sources, charges=[dict(
        charge_id='c', source_id='a', document_id='INV-1', reference_no='INV-1', amount=100,
        currency=None, charge_item='THC', evidence=[dict(source_id='a', anchor='p1', quote='INV-1')])]))
    bundle = result['material_bundle']
    assert {s['source_id'] for s in bundle['shared_context']} == {'a', 'b'}
    assert bundle['document_facts'][0]['reference_no'] == 'INV-1'
    assert 'currency' in bundle['document_facts'][0]['missing_fields']
    assert bundle['document_facts'][0]['evidence'][0]['source_id'] == 'a'


def test_two_invoice_numbers_cannot_silently_share_one_logical_document():
    result = skill.validate_extraction(dict(result=dict(charges=[dict(
        charge_id=f'c{i}', source_id=f's{i}', document_id='wrongly-merged', reference_no=ref,
        evidence=[dict(source_id=f's{i}', anchor='text', quote=ref)])
        for i, ref in enumerate(('INV-1', 'INV-2'))])))
    assert any(i['code'] == 'DOCUMENT_IDENTITY_CONFLICT' for i in result['issues'])


def test_non_charge_document_headers_are_saved_for_other_batches():
    candidate = dict(documents=[dict(source_id='hbl', document_id='hbl-1', document_type='bill_of_lading',
        reference_no='HBL-1', bill_of_lading='HBL-1', container_no='CONT-1',
        evidence=[dict(source_id='hbl', quote='HBL-1')])])
    result = skill.validate_extraction(dict(result=candidate))
    assert result['documents'][0]['bill_of_lading'] == 'HBL-1'
    bundle = skill.prepare_extraction(dict(documents=result['documents'], sources=[]))['material_bundle']
    assert bundle['source_facts'][0]['container_no'] == 'CONT-1'


def test_plain_pdf_transport_identifier_requires_layout_role_evidence():
    source = dict(source_id='pdf', format='pdf', blocks=[dict(anchor='page:1/embedded', page=1,
        kind='pdf_embedded_text', text='BOOKING B/L\nBOOK-1 BILL-2')])
    document = dict(source_id='pdf', document_id='bill', bill_of_lading='BOOK-1', evidence=[dict(
        source_id='pdf', anchor='page:1/embedded', field='bill_of_lading', quote='BOOK-1')])
    result = skill.validate_extraction(dict(result=dict(documents=[document]), sources=[source]))
    assert any(i['code']=='PDF_IDENTITY_LAYOUT_REQUIRED' for i in result['issues'])
    source['blocks'].append(dict(anchor='page:1/layout', page=1, kind='pdf_layout_text',
                                 text='BOOKING      B/L\nBOOK-1       BILL-2'))
    document['bill_of_lading'] = 'BILL-2'
    document['evidence'].append(dict(source_id='pdf', anchor='page:1/layout',
                                    field='bill_of_lading', quote='BOOKING      B/L\nBOOK-1       BILL-2'))
    result = skill.validate_extraction(dict(result=dict(documents=[document]), sources=[source]))
    assert not any(i['code']=='PDF_IDENTITY_LAYOUT_REQUIRED' for i in result['issues'])


def test_layout_repair_does_not_retain_old_computed_role_error():
    spec = importlib.util.spec_from_file_location('_layout_review',
        Path(__file__).parents[2] / 'human-review/scripts/skill.py')
    review = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(review)
    old = dict(documents=[dict(source_id='pdf', document_id='doc', bill_of_lading='BOOK-1')],
        issues=[dict(code='PDF_IDENTITY_LAYOUT_REQUIRED', document_id='doc', field='bill_of_lading')])
    combined = review.merge_extraction(dict(task=dict(revision=1, charges=[], extraction=old),
        result=dict(documents=[dict(source_id='pdf', document_id='doc', bill_of_lading='BILL-2')], issues=[])))
    assert combined['issues'] == []
    assert combined['documents'][0]['bill_of_lading'] == 'BILL-2'


def test_email_bundle_is_shared_across_batches_without_borrowing_document_identity():
    sources = [dict(source_id='mail', format='eml', name='request.eml', blocks=[
        dict(anchor='mail/body', kind='mail_body', text='Invoice INV-9. Carrier charges Customer handling USD 50.')]),
        dict(source_id='pdf', format='pdf', parent_source_id='mail', document_id='invoice-9',
             blocks=[dict(anchor='p1', text='INV-9 THC USD 100')]),
        dict(source_id='jpg', format='image', parent_source_id='mail', document_id='separate-note',
             blocks=[dict(anchor='frame1', text='EUR 30')])]
    result = skill.prepare_extraction(dict(id='bill', sources=sources))
    bundle = result['material_bundle']
    assert bundle['scope'] == 'submission' and bundle['bundle_id'] == 'bill'
    assert len(bundle['sources']) == 3
    assert bundle['sources'][1]['parent_source_id'] == 'mail'
    assert bundle['shared_context'][0]['source_id'] == 'mail'
    assert bundle['shared_context'][0]['anchor'] == 'mail/body'
    assert result['chunks'][1]['document_id'] == 'invoice-9'
    assert result['chunks'][2]['document_id'] == 'separate-note'
    assert result['chunks'][1]['context_after'] == ''
    assert bundle['shared_context'][0]['chunk_ids'] == [result['chunks'][0]['chunk_id']]


def test_shared_mail_context_is_bounded_and_omitted_text_still_has_owned_chunks():
    sources = [dict(source_id=f'mail-{i}', format='eml', blocks=[dict(anchor='mail/body',
        kind='mail_body', text='X' * 3000)]) for i in range(6)]
    result = skill.prepare_extraction(dict(sources=sources))
    shared = result['material_bundle']['shared_context']
    assert sum(len(item['text']) for item in shared) == 6000
    assert all(item['truncated'] for item in shared)
    assert all(item['chunk_ids'] for item in shared)
    assert sum(len(chunk['text']) for chunk in result['chunks']) == 18000


def test_untrusted_preview_never_enters_shared_mail_evidence():
    result = skill.prepare_extraction(dict(sources=[dict(source_id='mail', format='eml', blocks=[
        dict(anchor='native', kind='mail_body', authoritative=False, text='USD 100'),
        dict(anchor='body', kind='mail_body', text='handling $50')])]))
    assert [item['anchor'] for item in result['material_bundle']['shared_context']] == ['body']


def test_native_chunks_own_every_character_once_and_have_boundary_context():
    text = "THC USD 100.00\n" * 100
    payload = {"sources": [{"source_id": "d", "document_id": "logical-1", "blocks": [
        {"anchor": "p1", "page": 1, "text": text}, {"anchor": "p2", "page": 2, "text": "continuation USD 20"}]}],
        "chunk_size": 256, "context_chars": 50}
    result = skill.prepare_extraction(payload)
    assert result["valid"]
    page1 = [c for c in result["chunks"] if c["anchor"] == "p1"]
    assert "".join(c["text"] for c in page1) == text
    assert page1[-1]["context_after"] == "continuation USD 20"
    assert page1[-1]["context_after_source"]["anchor"] == "p2"
    assert all(c["owned_range"]["end"] - c["owned_range"]["start"] == len(c["text"]) for c in page1)


def test_different_documents_never_share_boundary_context_and_omissions_are_visible():
    sources = [{"source_id": "a", "text": "A" * 1000}, {"source_id": "b", "text": "B" * 1000}]
    result = skill.prepare_extraction({"sources": sources, "chunk_size": 256, "context_chars": 20, "max_chunks": 1})
    assert not result["valid"] and result["omitted"] > 0
    full = skill.prepare_extraction({"sources": sources, "chunk_size": 256, "context_chars": 20})
    first_b = next(c for c in full["chunks"] if c["source_id"] == "b")
    assert first_b["context_before"] == ""


def test_coverage_reports_unread_sources_instead_of_a_chunk_checklist():
    payload = {"result": {"charges": []}, "sources": [{"source_id": "d", "status": "partial"}]}
    assert not skill.validate_extraction(payload)["valid"]
    payload["sources"][0]["status"] = "ready"
    assert skill.validate_extraction(payload)["valid"]


@pytest.mark.parametrize("value", [True, "NaN", "Infinity"])
def test_native_structured_schema_rejects_invalid_financial_values(value):
    result = skill.validate_extraction({"result": {"charges": [{"amount": value}]}})
    assert not result["valid"] and result["issues"][0]["code"] == "EXTRACTION_SCHEMA"


def test_ocr_waiting_and_preview_do_not_become_empty_success():
    sources = [{"source_id": "pdf", "status": "waiting_url", "metadata": {"ocr_ranges": [{"pages": [1, 2], "status": "waiting_url"}]},
                "blocks": [{"source_id": "pdf", "anchor": "native:1", "text": "THC USD 100", "authoritative": False}]}]
    result = skill.prepare_extraction({"sources": sources})
    assert not result["valid"] and result["chunks"] == []
    assert result["issues"][0]["code"] == "SOURCE_INCOMPLETE"
    assert not skill.validate_coverage({"sources": sources, "coverage": []})["valid"]


def test_word_structure_row_is_context_only_and_not_reextracted():
    result = skill.prepare_extraction({"sources": [{"source_id": "word", "blocks": [
        {"anchor": "table:1/row:1", "text": "THC | USD | 100", "structure_only": True},
        {"anchor": "table:1/row:1/cell:1/paragraph:1", "text": "THC USD 100"}]}]})
    assert result["valid"] and len(result["chunks"]) == 1
    assert result["chunks"][0]["anchor"] == "table:1/row:1/cell:1/paragraph:1"


def test_native_schema_exposes_precise_references_and_preserves_real_locations():
    from langchain.agents.structured_output import ToolStrategy
    from langchain.tools import tool

    @tool
    def save_extraction(result: skill.ExtractionResult) -> dict:
        """Validate a native extraction result without a separate parser."""
        return result.model_dump(mode="python")

    expected = {"field", "line_id", "quote_start", "quote_end", "sheet", "row", "column", "cell",
                "coordinate", "cell_anchor", "paragraph", "table", "page"}
    for schema in (skill.extraction_schema({}), save_extraction.args_schema.model_json_schema(),
                   ToolStrategy(skill.ExtractionResult).schema_specs[0].json_schema):
        props = schema["$defs"]["EvidenceReference"]["properties"]
        assert expected <= props.keys()
        assert all(props[key].get("description") for key in expected)
        assert not {"verified", "content_sha256"} & props.keys()

    ref = {"source_id": "excel", "anchor": "sheet:1/row:4", "quote": "D4=20", "field": "amount",
           "line_id": "row:4", "quote_start": 18, "quote_end": 23, "sheet": "Charges",
           "row": 4, "column": "D", "cell": "D4", "coordinate": "D4", "cell_anchor": "sheet:1/D4"}
    result = save_extraction.invoke({"result": {"charges": [{"amount": "20", "evidence": [ref]}]}})
    assert result["charges"][0]["evidence"] == [ref]
    assert result["charges"][0]["raw_charge_item"] is None


def test_optional_null_offsets_dump_as_absent_but_business_nulls_remain():
    result = skill.ExtractionResult.model_validate({"charges": [{"evidence": [
        {"source_id": "word", "anchor": "body:1", "quote": "Amount USD 20", "quote_start": None,
         "quote_end": None, "field": None, "line_id": None}]}]}).model_dump(mode="python")
    assert result["charges"][0]["evidence"] == [
        {"source_id": "word", "anchor": "body:1", "quote": "Amount USD 20"}]
    assert result["charges"][0]["creditor"] is None


@pytest.mark.parametrize("offset", [-1, True, "4"])
def test_structured_quote_offsets_require_actual_nonnegative_integer(offset):
    result = skill.validate_extraction({"charges": [{"evidence": [
        {"source_id": "word", "quote": "USD 20", "quote_start": offset}]}]})
    assert not result["valid"] and result["issues"][0]["code"] == "EXTRACTION_SCHEMA"


def test_linked_word_cell_metadata_survives_human_revision_and_native_schema(tmp_path):
    from docx import Document

    modules = {}
    for name in ("document-intake", "evidence-linking", "billing-reconciliation", "human-review"):
        spec = importlib.util.spec_from_file_location("_word_revision_" + name.replace("-", "_"),
            Path(__file__).parents[2] / name / "scripts" / "skill.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        modules[name] = module
    document = Document()
    document.add_paragraph("Invoice INV-1 | Carrier bills Client")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Doc Fee | Amount USD 20"
    path = tmp_path / "invoice.docx"
    document.save(path)
    parsed = modules["document-intake"].intake_document(path, output_dir=tmp_path / "parsed", office_preview=False)
    assert parsed["status"] == "ready"
    source = parsed["sources"][0]
    block = next(row for row in source["blocks"] if row.get("kind") == "paragraph" and "Doc Fee" in row["text"])
    assert block["cell"] == 1 and block["row"] == 1 and isinstance(block["paragraph"], int)
    assert isinstance(block["table"], str)
    prepared = skill.prepare_extraction({"sources": parsed["sources"]})
    charge = {"charge_id": "c1", "source_id": source["source_id"], "line_id": block["anchor"],
              "charge_item": "Doc Fee", "amount": "25", "currency": "USD", "reference_no": "INV-1",
              "creditor": "Carrier", "debtor": "Client",
              "evidence": [{"source_id": source["source_id"], "anchor": block["anchor"], "quote": block["text"]}]}
    coverage = [{"source_id": chunk["source_id"], "anchor": chunk["anchor"], "chunk_id": chunk["chunk_id"],
                 "status": "extracted" if chunk["anchor"] == block["anchor"] else "no_charge", "reason": "checked",
                 "charge_ids": ["c1"] if chunk["anchor"] == block["anchor"] else []} for chunk in prepared["chunks"]]
    extraction = {"charges": [charge], "coverage": coverage}
    assert skill.validate_extraction({"result": extraction, "chunks": prepared["chunks"]})["valid"]
    task = {"revision": 1, "accounting_entity": "Client", "existing_records": [], "sources": parsed["sources"],
            "coverage": coverage, "chunks": prepared["chunks"], "extraction": extraction, "charges": [charge]}
    task["charges"] = modules["evidence-linking"].link_evidence(task)["charges"]
    assert task["charges"][0]["evidence"][0]["cell"] == 1
    task["records"] = modules["billing-reconciliation"].reconcile_charges(task)["records"]
    review = modules["human-review"].apply_revision(dict(task, expected_revision=1, reason="核对原始费用行",
        operations=[{"op": "set", "charge_id": "c1", "field": "amount", "value": "20"}]))
    assert review["valid"]
    current = modules["human-review"].prepare_recalculation({"task": task, "review": review})
    checked = skill.validate_extraction({"result": current["extraction"], "chunks": current["chunks"],
                                          "sources": current["sources"]})
    assert checked["valid"], checked["issues"]
    assert checked["charges"][0]["amount"] == 20
    assert checked["charges"][0]["evidence"][0]["cell"] == 1
    assert current["revision"] == 2 and task["revision"] == 1 and task["charges"][0]["amount"] == "25"


def test_global_raw_values_are_declared_in_native_schema_and_preserved_verbatim():
    from decimal import Decimal
    from langchain.agents.structured_output import ToolStrategy

    schema = ToolStrategy(skill.ExtractionResult).schema_specs[0].json_schema
    raw_fields = {"raw_amount", "raw_currency", "raw_quantity", "raw_unit_price"}
    charge_properties = schema["$defs"]["ChargeCandidate"]["properties"]
    assert raw_fields <= charge_properties.keys()
    assert all(charge_properties[field].get("description") for field in raw_fields)
    assert {"raw_amount", "raw_currency"} <= schema["$defs"]["DocumentTotal"]["properties"].keys()
    original = {"raw_charge_item": "Frais de dossier", "charge_item": "Frais de dossier",
                "raw_amount": "1\u202f234,56", "amount": "1234.56", "raw_currency": "€", "currency": "EUR",
                "raw_quantity": "2", "quantity": "2", "raw_unit_price": "617,28", "unit_price": "617.28"}
    result = skill.ExtractionResult.model_validate({"charges": [original], "totals": [
        {"role": "total", "raw_amount": "1.234,56", "amount": "1234.56", "raw_currency": "EUR", "currency": "EUR"}
    ]}).model_dump(mode="python")
    for field in raw_fields | {"raw_charge_item"}:
        assert result["charges"][0][field] == original[field]
    assert result["charges"][0]["amount"] == Decimal("1234.56")
    assert result["charges"][0]["reference_no"] is None
    assert result["totals"][0]["raw_amount"] == "1.234,56"


@pytest.mark.parametrize("field", ["raw_amount", "raw_currency", "raw_quantity", "raw_unit_price"])
def test_raw_financial_values_must_remain_source_text(field):
    result = skill.validate_extraction({"result": {"charges": [{field: 1234}]}})
    assert not result["valid"]
    assert any(issue["code"] == "EXTRACTION_SCHEMA" and issue["field"].endswith(field)
               for issue in result["issues"])


def test_unresolved_global_number_and_currency_survive_as_null_and_explicit_issue():
    result = skill.validate_extraction({"result": {"charges": [{
        "charge_item": "搬入料", "raw_amount": "1,234", "amount": None,
        "raw_currency": "$", "currency": None, "quantity": None, "unit_price": None,
        "evidence": [{"source_id": "bill", "quote": "搬入料 $ 1,234", "field": "amount"}]}],
        "issues": [{"code": "NUMBER_FORMAT_AMBIGUOUS", "field": "amount", "severity": "error",
                    "message": "1,234 may mean 1234 or 1.234; dollar currency not identified"}],
        "coverage": [{"source_id": "bill", "status": "needs_review", "reason": "ambiguous amount/currency"}]}})
    charge = result["charges"][0]
    assert not result["valid"] and charge["amount"] is None and charge["currency"] is None
    assert charge["raw_amount"] == "1,234" and charge["raw_currency"] == "$"
    assert charge["quantity"] is None and charge["unit_price"] is None
    assert any(issue["code"] == "NUMBER_FORMAT_AMBIGUOUS" for issue in result["issues"])


@pytest.mark.parametrize("value", [True, "NaN", "-Infinity", "1.234,56"])
def test_total_canonical_numbers_are_finite_and_not_locale_parsed(value):
    result = skill.validate_extraction({"result": {"totals": [{"role": "total", "amount": value}]}})
    assert not result["valid"]
    assert any(issue["code"] == "EXTRACTION_SCHEMA" and issue["field"] == "totals.0.amount"
               for issue in result["issues"])


def test_pre_alert_reference_is_not_inferred_from_the_subject_label():
    source = dict(source_id='mail', format='eml', status='ready', blocks=[
        dict(anchor='mail/headers', kind='mail_headers',
             text='Subject: 深圳容宇 PRE-ALERT SSZSE26093128 VASINGER LLC'),
        dict(anchor='mail/body', kind='mail_body', text='SHA pay TAR：H/C $50')])
    payload = {'result': {'charges': [{
        'source_id': 'mail', 'charge_item': 'H/C', 'amount': '50',
        'evidence': [{'source_id': 'mail', 'anchor': 'mail/body', 'quote': 'SHA pay TAR：H/C $50'}]}]},
        'sources': [source]}

    charge = skill.validate_extraction(payload)['charges'][0]
    assert charge['reference_no'] is None
    assert not [ref for ref in charge['evidence'] if ref.get('field') == 'reference_no']

    provided = {'result': {'charges': [{
        'source_id': 'mail', 'charge_item': 'H/C', 'amount': '50', 'reference_no': 'INV-9',
        'evidence': [{'source_id': 'mail', 'anchor': 'mail/body', 'quote': 'SHA pay TAR：H/C $50'},
                     {'source_id': 'mail', 'anchor': 'mail/headers', 'quote': 'Invoice No: INV-9',
                      'field': 'reference_no'}]}]}, 'sources': [source]}
    assert skill.validate_extraction(provided)['charges'][0]['reference_no'] == 'INV-9'


def test_email_declared_mbl_conflicting_with_attachment_b_l_is_recorded():
    sources = [
        dict(source_id='mail', format='eml', blocks=[dict(anchor='mail/headers', kind='mail_headers',
             text='Subject: 深圳容宇 PRE-ALERT SSZSE26093128 MBL NO._ OOLU4055331281')]),
        dict(source_id='mbl', format='pdf', parent_source_id='mail', blocks=[
             dict(anchor='page:1/embedded', kind='pdf_embedded_text',
                  text='B/L NO. HZMHK331281\nBOOKING 4055331281')])]
    result = skill.validate_extraction({'result': {'documents': [dict(
        source_id='mbl', document_id='d1', bill_of_lading='HZMHK331281',
        evidence=[dict(source_id='mbl', anchor='page:1/layout', field='bill_of_lading',
                       quote='HZMHK331281')])]}, 'sources': sources})
    assert any(issue['code'] == 'IDENTITY_CONFLICT' and 'OOLU4055331281' in issue['message']
               for issue in result['issues'])


def _representation_sources():
    layout_text = '\n'.join([
        'HANDLING CHARGE - 65.00 1.00 C 65.00',
        "4 ADD'L DETENTION FREE TIME CHARGE - 1,310.40 1.00 C 1,310.40",
        'HANDLING CHARGE - 35.00 1.00 C 35.00',
    ])
    layout_rows = [
        dict(description='HANDLING CHARGE', amount='65.00', raw_amount='65.00', currency='USD', line=1),
        dict(description="4 ADD'L DETENTION FREE TIME CHARGE", amount='1310.40', raw_amount='1,310.40', currency='USD', line=2),
        dict(description='HANDLING CHARGE', amount='35.00', raw_amount='35.00', currency='USD', line=3),
    ]
    ocr_text = json.dumps({'chargeList': [
        dict(chargeItem='HANDLING CHARGE', amount='100.00'),
        dict(chargeItem="4 ADD'L DETENTION FREE TIME CHARGE", amount='1310.40'),
    ]})
    return [dict(source_id='pdf', format='pdf', status='ready', blocks=[
        dict(anchor='page:1/layout/financial', kind='pdf_layout_financial',
             text=layout_text, rows=layout_rows, page=1),
        dict(anchor='ocr/pages:1-1/document:1', kind='ocr_structured', text=ocr_text),
    ])]


def test_ocr_merged_row_conflicts_with_layout_and_candidate_uses_layout():
    sources = _representation_sources()
    charges = [
        dict(charge_id='c65', source_id='pdf', charge_item='HANDLING CHARGE', amount=65,
             currency='USD', evidence=[dict(source_id='pdf', anchor='page:1/layout/financial',
                                            field=field) for field in ('charge_item', 'amount', 'currency')]),
        dict(charge_id='c1310', source_id='pdf', charge_item="4 ADD'L DETENTION FREE TIME CHARGE",
             amount='1310.40', currency='USD', evidence=[dict(source_id='pdf', anchor='page:1/layout/financial',
                                                               field=field) for field in ('charge_item', 'amount', 'currency')]),
        dict(charge_id='c35', source_id='pdf', charge_item='HANDLING CHARGE', amount=35,
             currency='USD', evidence=[dict(source_id='pdf', anchor='page:1/layout/financial',
                                            field=field) for field in ('charge_item', 'amount', 'currency')]),
    ]
    result = skill.validate_extraction({'result': {'charges': charges}, 'sources': sources})
    assert result['valid']
    assert not any(issue.get('severity') == 'error' for issue in result['issues'])
    assert any(issue['code'] == 'OCR_LAYOUT_CONFLICT' and '合并' in issue['message']
               for issue in result['issues'])
    assert not any(issue['code'] == 'PDF_LAYOUT_EVIDENCE_REQUIRED' for issue in result['issues'])


def test_layout_blocks_a_candidate_amount_only_supported_by_ocr():
    sources = _representation_sources()
    charge = dict(charge_id='c100', source_id='pdf', charge_item='HANDLING CHARGE',
                  amount=100, currency='USD', evidence=[dict(
                      source_id='pdf', anchor='ocr/pages:1-1/document:1',
                      kind='ocr_structured', field=field)
                      for field in ('charge_item', 'amount', 'currency')])
    result = skill.validate_extraction({'result': {'charges': [charge]}, 'sources': sources})
    assert not result['valid']
    assert any(issue['code'] == 'OCR_LAYOUT_CONFLICT' and issue.get('severity') == 'error'
               for issue in result['issues'])
    assert any(issue['code'] == 'PDF_LAYOUT_EVIDENCE_REQUIRED' for issue in result['issues'])
