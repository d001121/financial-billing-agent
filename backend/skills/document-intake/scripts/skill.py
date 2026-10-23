"""Portable document intake. Format parsing uses mature libraries, not an agent runner."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import mimetypes
import posixpath
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import importlib.util
from contextlib import nullcontext
from time import perf_counter
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
from threading import RLock
from datetime import date, datetime, time, timedelta, timezone
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any, Callable, Iterator
from urllib.parse import urlparse

from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document


VERSION = "1.8.0"
RAW_OCR_CACHE = True
IMAGE_TEXT_OCR = True
PDF_TEXT_MODES = {'full', 'ocr_only', 'markdown'}
ATTACHMENT_CONCURRENCY = True
MARKDOWN_FIRST = True
OCR_ENDPOINT = "http://36.103.199.11:5002/process"

_LAYOUT_CURRENCIES = frozenset({
    'AED', 'AUD', 'BRL', 'CAD', 'CHF', 'CNY', 'DKK', 'EUR', 'GBP', 'HKD',
    'IDR', 'INR', 'JPY', 'KRW', 'MXN', 'MYR', 'NOK', 'NZD', 'PHP', 'PLN',
    'RMB', 'SAR', 'SEK', 'SGD', 'THB', 'TRY', 'TWD', 'USD', 'VND', 'ZAR',
})
_LAYOUT_TOTAL = re.compile(
    r'^(?:TOTAL|GRAND\s+TOTAL|SUB\s*TOTAL|PAID\s+AMOUNT|BALANCE|AMOUNT\s+DUE|TOTAL\s+CHARGES)\b',
    re.IGNORECASE,
)
_LAYOUT_ROW = re.compile(
    r'^(?P<description>.+?)\s+(?P<unit>[A-Za-z*/-]+)\s+'
    r'(?P<rate>\(?[+-]?\d[\d,]*(?:\.\d+)?\)?)\s+'
    r'(?P<qty>\d+(?:\.\d+)?)\s+(?P<pc>[A-Z])\s+'
    r'(?P<amount>\(?[+-]?\d[\d,]*(?:\.\d+)?\)?)\s*$'
)
_LAYOUT_CURRENCY_ROW = re.compile(
    r'^(?P<description>.+?)\s+(?P<currency>[A-Z]{3})\s+'
    r'(?P<amount>\(?[+-]?\d[\d,]*(?:\.\d+)?\)?)\s*(?P<trailing>[A-Z])?\s*$'
)
_LAYOUT_DESCRIPTION_PREFIX = re.compile(
    r'^(?:[A-Z]{4}\s+)?[A-Z0-9]{6,}\s+Waybill\s+(?:[A-Z0-9]+\s+)?', re.IGNORECASE
)


def _clean_layout_description(value):
    return _LAYOUT_DESCRIPTION_PREFIX.sub('', str(value or ''), count=1).strip()


def _layout_decimal(value):
    raw = str(value).strip()
    negative = raw.startswith('(') and raw.endswith(')')
    try:
        number = Decimal(raw.strip('()').replace(',', '').lstrip('+'))
    except InvalidOperation:
        return None
    if not number.is_finite():
        return None
    return -number if negative else number


def layout_financial_rows(text: str) -> list[dict]:
    """Parse charge-like rows from layout-preserved PDF text.

    This is deliberately conservative: it keeps the original line as evidence
    and only recognises an ISO currency or the common rate/qty/P-C/amount
    table shape. Totals and prose are not treated as charge rows.
    """
    source_text = str(text or '')
    default_currency = None
    match = re.search(r'\bCURRENCY\s+([A-Z]{3})\b', source_text)
    if match and match.group(1) in _LAYOUT_CURRENCIES:
        default_currency = match.group(1)
    rows, seen = [], set()
    for line_number, raw in enumerate(source_text.splitlines(), 1):
        line = raw.strip()
        if not line or not re.search(r'[A-Za-z]', line) or not re.search(r'\d', line):
            continue
        if _LAYOUT_TOTAL.match(line):
            continue
        row = None
        match = _LAYOUT_ROW.match(line)
        if match:
            amount = _layout_decimal(match.group('amount'))
            rate = _layout_decimal(match.group('rate'))
            quantity = _layout_decimal(match.group('qty'))
            if amount is not None and rate is not None and quantity is not None:
                row = {
                    'description': _clean_layout_description(match.group('description')),
                    'currency': default_currency,
                    'raw_amount': match.group('amount'),
                    'amount': format(amount, 'f'),
                    'raw_rate': match.group('rate'),
                    'rate': format(rate, 'f'),
                    'raw_quantity': match.group('qty'),
                    'quantity': format(quantity, 'f'),
                    'pc': match.group('pc'),
                    'line': line_number,
                    'text': line,
                }
        if row is None:
            match = _LAYOUT_CURRENCY_ROW.match(line)
            if match and match.group('currency') in _LAYOUT_CURRENCIES:
                amount = _layout_decimal(match.group('amount'))
                if amount is not None:
                    row = {
                        'description': _clean_layout_description(match.group('description')),
                        'currency': match.group('currency'),
                        'raw_amount': match.group('amount'),
                        'amount': format(amount, 'f'),
                        'raw_rate': None,
                        'rate': None,
                        'raw_quantity': None,
                        'quantity': None,
                        'pc': match.group('trailing'),
                        'line': line_number,
                        'text': line,
                    }
        if row is None:
            continue
        key = (row['description'].casefold(), row['amount'], row['currency'])
        if key in seen:
            continue
        seen.add(key)
        rows.append(row)
    return rows

_md_spec = importlib.util.spec_from_file_location(__name__ + '_markdown', Path(__file__).with_name('pdf_ocr_md.py'))
_md = importlib.util.module_from_spec(_md_spec)
_md_spec.loader.exec_module(_md)

_cache_spec = importlib.util.spec_from_file_location(__name__ + '_ocr_cache', Path(__file__).with_name('ocr_cache.py'))
_cache = importlib.util.module_from_spec(_cache_spec)
_cache_spec.loader.exec_module(_cache)


def write_markdown(source: dict, output_dir: str | Path) -> dict:
    """Portable package entry also used after an explicit PDF supplement."""
    source['markdown'] = _md.build_markdown(source, output_dir)
    return source['markdown']


def supplement_ocr(source: dict, *, path: str | Path, output_dir: str | Path,
                   reason: str, pages: list[int] | None = None, frame: int | None = None,
                   ocr_endpoint: str = OCR_ENDPOINT, ocr_request: Callable | None = None,
                   url_resolver: Callable | None = None, ocr_text_pointer: str | None = None,
                   obs_options: dict | None = None, request_guard: Callable | None = None,
                   ocr_response_format: str = 'text', ocr_cache_dir=None,
                   ocr_cache_ttl=86400, ocr_cache_version='orderChange-v1', on_progress=None) -> dict:
    """Read selected original PDF pages/image frame through the configured OCR.

    Reuses intake's publisher, wire-response storage and verified response mapping.
    No image/model call or inferred physical-page coordinates are produced.
    """
    source, path, output = deepcopy(source), Path(path), Path(output_dir)
    if not str(reason).strip() or source.get('format') not in {'pdf', 'image'}:
        raise ValueError('OCR补读需要PDF或图片来源及具体原因')
    if 'sha256:' + _digest(path.read_bytes()) != source.get('content_id'):
        raise ValueError('原件内容已变化，不能绑定旧来源')
    if source['format'] == 'pdf':
        from pypdf import PdfReader, PdfWriter
        if frame is not None:
            raise ValueError('PDF使用pages指定原页，不使用frame')
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF需要密码')
        if not pages or any(type(p) is not int or not 1 <= p <= len(reader.pages) for p in pages):
            raise ValueError('pages必须指定原PDF真实页范围内的页号')
        targets = list(dict.fromkeys(pages))
    else:
        from PIL import Image
        if pages is not None:
            raise ValueError('图片使用frame指定原帧，不使用pages')
        frame = 1 if frame is None else frame
        with Image.open(path) as image:
            count = getattr(image, 'n_frames', 1)
        if type(frame) is not int or not 1 <= frame <= count:
            raise ValueError('frame必须在原图片真实帧范围内')
        targets = [frame]
    worker = Intake(output, ocr_endpoint=ocr_endpoint, ocr_request=ocr_request,
        url_resolver=url_resolver, ocr_text_pointer=ocr_text_pointer,
        pdf_pages_per_slice=1, max_depth=0, max_files=1, max_total_bytes=max(1, path.stat().st_size),
        doc_converter=None, obs_options=obs_options, request_guard=request_guard,
        ocr_response_format=ocr_response_format, office_preview=False, pdf_text_mode='ocr_only',
        ocr_cache_dir=ocr_cache_dir, ocr_cache_ttl=ocr_cache_ttl,
        ocr_cache_version=ocr_cache_version, on_progress=on_progress)
    source.setdefault('metadata', {})
    before = len(source['blocks'])
    reads = []
    attempt = uuid4().hex
    for target in targets:
        # A successful explicit recheck is immutable and reused within this task.
        previous = next((r for r in reversed(source['metadata'].get('ocr_rechecks', []))
            if r.get('kind') == source['format'] and r.get('target') == target
            and r.get('status') == 'ready' and r.get('source_content_id') == source['content_id']), None)
        if previous:
            reads.append(dict(previous, reused=True))
            continue
        if source['format'] == 'pdf':
            selected = output / 'pdf-slices' / source['source_id'] / f'page-{target}.pdf'
            selected.parent.mkdir(parents=True, exist_ok=True)
            writer = PdfWriter()
            writer.add_page(reader.pages[target - 1])
            writer.write(selected)
            mapping = dict(source_id=source['source_id'], original_pages=[target], kind='pdf_slice')
            worker.artifact(selected, 'pdf_slice', source, original_pages=[target])
            anchor = f'ocr/recheck:{attempt}/page:{target}'
        else:
            selected = path
            if count > 1:
                selected = output / 'image-frames' / source['source_id'] / f'frame-{target}.png'
                selected.parent.mkdir(parents=True, exist_ok=True)
                with Image.open(path) as image:
                    image.seek(target - 1)
                    image.convert('RGB').save(selected)
                worker.artifact(selected, 'image_frame', source, original_frame=target)
            mapping = dict(source_id=source['source_id'], frame=target, kind='image_frame')
            anchor = f'ocr/recheck:{attempt}/frame:{target}'
        status = worker.ocr(source, selected, None, mapping, anchor)
        receipt = dict(kind=source['format'], target=target, status=status, reason=reason,
            source_content_id=source['content_id'], anchor=anchor, method='specified_ocr',
            input_sha256=_digest(selected.read_bytes()))
        source['metadata'].setdefault('ocr_rechecks', []).append(receipt)
        reads.append(receipt)
    # Keep the original ranges/frames and status intact: a successful recheck of
    # one page does not prove a previously failed multi-page intake is complete.
    source['text'] = '\n'.join(b['text'] for b in source['blocks']
        if b.get('authoritative', True) and not b.get('structure_only'))
    worker.markdown(source)
    return dict(source=source, added_anchors=[b['anchor'] for b in source['blocks'][before:]],
        method='ocr', reads=reads, issues=worker.issues, artifacts=worker.artifacts)


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    return value


def _safe_name(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    if not name or name.split(".")[0].upper() in {"CON", "PRN", "AUX", "NUL", *[f"COM{i}" for i in range(1, 10)], *[f"LPT{i}" for i in range(1, 10)]}:
        name = "attachment_" + name
    return name[:180]


def _format(path: Path, data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith((b"\x89PNG", b"\xff\xd8\xff", b"GIF8", b"II*\x00", b"MM\x00*", b"II+\x00", b"MM\x00+", b"BM")):
        return "image"
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" in names:
                return "docx"
            if "xl/workbook.xml" in names:
                return "xlsx"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls" if path.suffix.lower() == ".xls" else "doc"
    return {".eml": "eml", ".json": "json", ".txt": "text", ".csv": "text", ".html": "html", ".htm": "html", ".doc": "doc", ".xls": "xls"}.get(path.suffix.lower(), "unsupported")


class OBSPublisher:
    """Publish OCR inputs using Huawei's SDK, without changing bucket/object ACLs."""

    def __init__(self, options: dict, before_request: Callable | None = None):
        from obs import ObsClient
        import certifi
        for name in ("access_key_id", "secret_access_key", "server", "bucket"):
            if not options.get(name):
                raise ValueError(f"OBS配置缺少{name}")
        if urlparse(options["server"]).scheme != "https":
            raise ValueError("OBS_SERVER需要HTTPS地址")
        self.bucket = options["bucket"]
        self.prefix = str(options.get("object_prefix", "images")).strip("/")
        self.expires = int(options.get("expires", 3600))
        timeout = int(options.get("socket_timeout", 60))
        if self.expires < 300 or timeout <= 0:
            raise ValueError("OBS签名有效期至少300秒，超时必须为正数")
        self.client = ObsClient(access_key_id=options["access_key_id"],
            secret_access_key=options["secret_access_key"],
            security_token=options.get("security_token") or None,
            server=options["server"], timeout=timeout,
            ssl_verify=certifi.where(), max_retry_count=0)
        self.before_request = before_request
        self.receipts: list[dict] = []

    def publish(self, path: Path, mapping: dict) -> str:
        from obs import PutObjectHeader
        with path.open("rb") as stream:
            digest = hashlib.file_digest(stream, "sha256").hexdigest()
        suffix = path.suffix.lower()
        if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
            suffix = ".bin"
        key = "/".join(part for part in (self.prefix, "workorder-agent", digest + suffix) if part)
        if self.before_request:
            self.before_request()
        try:
            response = self.client.putFile(self.bucket, key, str(path),
                metadata={"sha256": digest},
                headers=PutObjectHeader(contentType=mimetypes.guess_type(path.name)[0] or "application/octet-stream"))
        except Exception as exc:
            raise RuntimeError(f"OBS上传请求失败（{type(exc).__name__}）") from exc
        if not response.status or not 200 <= response.status < 300:
            raise RuntimeError(f"OBS上传失败：HTTP {response.status}，{response.errorCode or 'unknown'}，request_id={response.requestId}")
        receipt = dict(bucket=self.bucket, object_key=key, sha256=digest,
            size_bytes=path.stat().st_size, status=response.status,
            etag=getattr(response.body, "etag", None), request_id=response.requestId,
            uploaded_at=datetime.now(timezone.utc).isoformat(), **mapping)
        self.receipts.append(receipt)
        signed = self.client.createSignedUrl("GET", bucketName=self.bucket, objectKey=key, expires=self.expires)
        url = signed["signedUrl"]
        if urlparse(url).scheme != "https":
            raise ValueError("OBS未返回有效的HTTPS签名下载地址")
        receipt["url_expires_at"] = (datetime.now(timezone.utc) + timedelta(seconds=self.expires)).isoformat()
        return url

    def close(self):
        self.client.close()


class Intake:
    def __init__(self, output_dir: Path, *, ocr_endpoint: str, ocr_request: Callable | None, url_resolver: Callable | None, ocr_text_pointer: str | None, pdf_pages_per_slice: int, max_depth: int, max_files: int, max_total_bytes: int, doc_converter: str | None, obs_options: dict | None = None, request_guard: Callable | None = None, ocr_response_format: str = "text", office_preview: bool = True, pdf_text_mode: str = 'full', attachment_concurrency: int = 1, ocr_cache_dir=None, ocr_cache_ttl=86400, ocr_cache_version='orderChange-v1', on_progress=None, image_text_ocr: bool = False):
        if min(pdf_pages_per_slice, max_files, max_total_bytes) <= 0 or max_depth < 0:
            raise ValueError("Invalid intake limits")
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.ocr_endpoint, self.ocr_request = ocr_endpoint, ocr_request
        self.url_resolver, self.ocr_text_pointer = url_resolver, ocr_text_pointer
        if ocr_response_format not in {"text", "orderChange"}:
            raise ValueError("ocr_response_format必须为text或orderChange")
        self.ocr_response_format = ocr_response_format
        self.pdf_pages_per_slice = pdf_pages_per_slice
        self.max_depth, self.max_files, self.max_total_bytes = max_depth, max_files, max_total_bytes
        self.doc_converter = doc_converter
        self.office_preview = office_preview
        if pdf_text_mode not in PDF_TEXT_MODES:
            raise ValueError('pdf_text_mode必须为full、ocr_only或markdown')
        self.pdf_text_mode = pdf_text_mode
        self.converter_version: str | None = None
        self.request_guard = request_guard
        self.obs_options = obs_options or {}
        self.publisher: OBSPublisher | None = None
        self.sources: list[dict] = []
        self.issues: list[dict] = []
        self.artifacts: list[dict] = []
        self.total_bytes = 0
        if not 1 <= attachment_concurrency <= 16:
            raise ValueError('attachment_concurrency必须介于1和16')
        self.attachment_concurrency = attachment_concurrency
        self.lock = RLock()
        self.publisher_lock = RLock()
        self.ocr_cache = _cache.OCRResponseCache(ocr_cache_dir, ocr_cache_ttl) if ocr_cache_dir else None
        self.ocr_cache_version = ocr_cache_version
        self.on_progress = on_progress
        self.image_text_ocr = image_text_ocr
        self.image_parser = None
        self.image_parser_lock = RLock()

    def progress(self, source, stage, **details):
        if self.on_progress:
            self.on_progress(dict(stage=stage, source_id=source['source_id'],
                name=source.get('name'), **details))

    def issue(self, source: dict, code: str, message: str, *, scope: str | None = None, blocking: bool = True):
        item = {"code": code, "message": message, "source_id": source["source_id"], "scope": scope, "blocking": blocking}
        self.issues.append(item)
        if blocking:
            source["status"] = "partial"

    def block(self, source: dict, anchor: str, text: str, **metadata):
        block = {"source_id": source["source_id"], "anchor": anchor, "text": text, "authoritative": True, **metadata}
        source["blocks"].append(block)

    def artifact(self, path: Path, kind: str, source: dict, **metadata):
        item = {"path": str(path.resolve()), "kind": kind, "source_id": source["source_id"], "sha256": _digest(path.read_bytes()), **metadata}
        self.artifacts.append(item)
        return item

    def parse(self, path: Path, *, parent_id: str | None = None, occurrence: str = "root", depth: int = 0, source_url: str | None = None, original_name: str | None = None) -> dict:
        data = path.read_bytes()
        content_id = "sha256:" + _digest(data)
        source_id = "src_" + _digest(f"{parent_id or ''}|{occurrence}|{content_id}".encode())[:24]
        source = {"source_id": source_id, "content_id": content_id, "parent_source_id": parent_id, "occurrence_id": occurrence, "name": original_name or path.name, "path": str(path.resolve()), "size_bytes": len(data), "format": "unknown", "mime_type": mimetypes.guess_type(original_name or path.name)[0] or "application/octet-stream", "status": "ready", "blocks": [], "text": "", "preview_text": "", "metadata": {}}
        with self.lock:
            self.sources.append(source)
            self.total_bytes += len(data)
            exceeds = depth > self.max_depth or len(self.sources) > self.max_files or self.total_bytes > self.max_total_bytes
        if exceeds:
            source["status"] = "failed"
            self.issue(source, "intake_limit", "材料超过递归深度、文件数或总字节限制；本项未解析。")
            source["status"] = "failed"
            self.markdown(source)
            return source
        try:
            kind = _format(path, data)
            source["format"] = kind
            source["mime_type"] = mimetypes.guess_type(source["name"])[0] or "application/octet-stream"
            if kind == "eml":
                self.eml(source, data, depth)
            elif kind == "pdf":
                self.pdf(source, path, source_url)
            elif kind == "image":
                self.image(source, path, source_url)
            elif kind in {"xlsx", "xls"}:
                self.excel(source, path, kind)
            elif kind == "docx":
                self.docx(source, path, depth)
            elif kind == "doc":
                self.doc(source, path, depth)
            elif kind in {"text", "html", "json"}:
                text = data.decode("utf-8-sig")
                if kind == "json":
                    source["metadata"]["json_value"] = json.loads(text)
                elif kind == "html":
                    text = self.html_text(text)
                self.block(source, "body", text, kind="body")
            else:
                self.issue(source, "unsupported_format", "无法识别或不支持该材料格式。")
                source["status"] = "failed"
        except Exception as exc:
            # A per-file failure remains a failed occurrence; siblings continue.
            self.issue(source, "parse_failed", f"{type(exc).__name__}: {exc}")
            source["status"] = "failed"
        authoritative = [block for block in source["blocks"] if block["authoritative"] and not block.get("structure_only")]
        previews = [block for block in source["blocks"] if not block["authoritative"]]
        source["preview_text"] = "\n".join(block["text"] for block in previews)
        source["text"] = "\n".join(block["text"] for block in authoritative)
        self.markdown(source)
        return source

    def markdown(self, source: dict):
        md = write_markdown(source, self.output_dir)
        self.artifact(Path(md['path']), 'markdown', source)
        self.artifact(Path(md['map_path']), 'markdown_map', source)

    @staticmethod
    def html_text(text: str) -> str:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text("\n", strip=True)

    def eml(self, source: dict, data: bytes, depth: int):
        pending = []
        message = BytesParser(policy=policy.default).parsebytes(data)
        source["metadata"]["headers"] = {key: str(message.get(key, "")) for key in ["Subject", "From", "To", "Cc", "Date", "Message-ID", "In-Reply-To"]}
        self.block(source, "mail/headers", "\n".join(f"{key}: {value}" for key, value in source["metadata"]["headers"].items() if value), kind="mail_headers")
        selected = message.get_body(preferencelist=("plain", "html"))
        if selected is not None:
            try:
                text = selected.get_content()
            except (UnicodeError, LookupError) as exc:
                self.issue(source, "mail_body_encoding", str(exc), scope="body")
            else:
                self.block(source, "mail/body", self.html_text(text) if selected.get_content_type() == "text/html" else text, kind="mail_body", content_type=selected.get_content_type())
        for defect in message.defects:
            self.issue(source, "mime_defect", str(defect))

        def visit(part, part_path: str):
            if part.get_content_type() == "message/rfc822":
                payload = part.get_payload()
                nested = payload if isinstance(payload, list) else []
                for i, child in enumerate(nested):
                    save(child.as_bytes(policy=policy.default), part.get_filename() or f"forwarded-{i + 1}.eml", f"{part_path}/{i + 1}", part)
                if not nested:
                    self.issue(source, "nested_mail_invalid", "嵌套邮件没有可读消息。", scope=part_path)
                return
            if part.is_multipart():
                for i, child in enumerate(part.iter_parts(), 1):
                    visit(child, f"{part_path}/{i}")
                return
            if part is selected:
                return
            filename = part.get_filename()
            if not filename and part.get_content_maintype() == "text" and part.get_content_disposition() != "attachment":
                source["metadata"].setdefault("alternate_body_parts", []).append(part_path)
                return
            payload = part.get_payload(decode=True)
            if payload is None:
                self.issue(source, "attachment_payload_missing", "附件没有可解码内容。", scope=part_path)
                return
            name = filename or f"inline-{part_path.replace('/', '-')}{mimetypes.guess_extension(part.get_content_type()) or '.bin'}"
            save(payload, name, part_path, part)

        def save(payload: bytes, name: str, part_path: str, part):
            target = self.output_dir / "attachments" / source["source_id"] / f"{part_path.replace('/', '-')}-{_safe_name(name)}"
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(payload)
            pending.append((target, name, part_path, part))

        def parse_attachment(item):
            target, name, part_path, part = item
            child = self.parse(target, parent_id=source["source_id"], occurrence=part_path, depth=depth + 1, original_name=name)
            child["metadata"].update({"content_id_header": part.get("Content-ID"), "disposition": part.get_content_disposition(), "declared_content_type": part.get_content_type()})
            self.artifact(target, "mail_attachment", child)
            if child["status"] != "ready":
                self.issue(source, "attachment_incomplete", f"附件 {name} 未完整处理。", scope=child["source_id"])

        visit(message, "mime")
        # One bounded pool only at the root EML; nested messages stay serial.
        if depth == 0 and self.attachment_concurrency > 1 and len(pending) > 1:
            with ThreadPoolExecutor(max_workers=self.attachment_concurrency) as pool:
                list(pool.map(parse_attachment, pending))
        else:
            for item in pending:
                parse_attachment(item)

    def pdf(self, source: dict, path: Path, source_url: str | None):
        from pypdf import PdfReader, PdfWriter
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("PDF需要密码，尚未读取")
        count = len(reader.pages)
        source["metadata"]["page_count"] = count
        source["metadata"]["pages"] = []
        embedded = {}
        layouts = {}
        source['metadata']['pdf_text_mode'] = self.pdf_text_mode
        if self.pdf_text_mode in {'full', 'markdown'}:
            try:
                from langchain_community.document_loaders import PDFPlumberLoader
                for document in PDFPlumberLoader(path, text_kwargs={'layout': True}, dedupe=True).load():
                    layouts[document.metadata['page'] + 1] = '\n'.join(
                        line.rstrip() for line in document.page_content.splitlines() if line.strip())
            except Exception as exc:
                self.issue(source, 'pdf_layout_failed', str(exc), blocking=False)
        for index, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                text = ""
                self.issue(source, "native_preview_failed", str(exc), scope=f"page:{index}", blocking=False)
            source["metadata"]["pages"].append({"page": index, "width": float(page.mediabox.width), "height": float(page.mediabox.height)})
            self.block(source, f"page:{index}/native", text, page=index, kind="native_preview", authoritative=False)
            # pypdf layout mode can omit an entire transformed text stream
            # (reproduced on the original HBL). Keep the full plain extraction;
            # original PDF preview remains available for column interpretation.
            embedded[index] = text
        if count == 0:
            raise ValueError("PDF没有页面")
        preview = deepcopy(source)
        preview.update(status='processing', text='\n'.join(layouts.values()) or '\n'.join(embedded.values()))
        preview['preview_text'] = preview['text']
        preview['metadata']['local_preview'] = True
        preview['metadata']['financial_rows'] = [dict(page=page, **row)
            for page, text in layouts.items() for row in layout_financial_rows(text)]
        self.progress(source, 'local_ready', preview=preview)
        ranges = []
        for start in range(0, count, self.pdf_pages_per_slice):
            end = min(start + self.pdf_pages_per_slice, count)
            part_path = path
            if count > self.pdf_pages_per_slice:
                part_path = self.output_dir / "pdf-slices" / source["source_id"] / f"pages-{start + 1}-{end}.pdf"
                part_path.parent.mkdir(parents=True, exist_ok=True)
                writer = PdfWriter()
                for page in reader.pages[start:end]:
                    writer.add_page(page)
                with part_path.open("wb") as stream:
                    writer.write(stream)
                self.artifact(part_path, "pdf_slice", source, original_pages=list(range(start + 1, end + 1)))
            mapping = {"original_pages": list(range(start + 1, end + 1)), "kind": "pdf_slice", "source_id": source["source_id"]}
            for page_number in mapping["original_pages"]:
                financial_rows = layout_financial_rows(layouts.get(page_number, ''))
                if financial_rows:
                    self.block(source, f"page:{page_number}/layout/financial",
                        "\n".join(row["text"] for row in financial_rows),
                        page=page_number, kind="pdf_layout_financial",
                        extraction_method="pdfplumber.layout_charge_rows",
                        rows=financial_rows, priority=0,
                        source_content_id=source["content_id"])
            result = self.ocr(source, part_path, source_url if part_path == path else None, mapping, f"ocr/pages:{start + 1}-{end}")
            ranges.append({"pages": mapping["original_pages"], "status": result})
            # The specified OCR remains mandatory. Its structured response is
            # one representation, not the complete searchable PDF. Preserve
            # embedded text separately, with the exact original page/hash.
            if result == "ready" and self.pdf_text_mode in {'full', 'markdown'}:
                for page_number in mapping["original_pages"]:
                    if embedded.get(page_number, "").strip():
                        self.block(source, f"page:{page_number}/embedded", embedded[page_number],
                            page=page_number, kind="pdf_embedded_text", extraction_method="pypdf.plain",
                            priority=20,
                            source_content_id=source["content_id"], ocr_status=result,
                            ocr_anchor=f"ocr/pages:{start + 1}-{end}")
                    if layouts.get(page_number, '').strip():
                        self.block(source, f'page:{page_number}/layout', layouts[page_number],
                            page=page_number, kind='pdf_layout_text',
                            extraction_method='langchain.PDFPlumberLoader(layout=True,dedupe=True)',
                            priority=10,
                            source_content_id=source['content_id'], ocr_status=result,
                            ocr_anchor=f'ocr/pages:{start + 1}-{end}')
        source["metadata"]["ocr_ranges"] = ranges
        if all(item["status"] == "waiting_url" for item in ranges):
            source["status"] = "waiting_url"

    def image(self, source: dict, path: Path, source_url: str | None):
        from PIL import Image
        with Image.open(path) as image:
            frames = getattr(image, "n_frames", 1)
            source["metadata"].update({"width": image.width, "height": image.height, "frame_count": frames, "image_format": image.format})
            results = []
            for frame in range(frames):
                image.seek(frame)
                image.load()
                part_path = path
                if frames > 1:
                    part_path = self.output_dir / "image-frames" / source["source_id"] / f"frame-{frame + 1}.png"
                    part_path.parent.mkdir(parents=True, exist_ok=True)
                    image.convert("RGB").save(part_path)
                    self.artifact(part_path, "image_frame", source, original_frame=frame + 1)
                results.append(self.ocr(source, part_path, source_url if frames == 1 else None, {"frame": frame + 1, "kind": "image_frame", "source_id": source["source_id"]}, f"ocr/frame:{frame + 1}"))
                if self.image_text_ocr:
                    self.image_text(source, image, frame + 1, results[-1])
            source["metadata"]["ocr_frames"] = [{"frame": i + 1, "status": status} for i, status in enumerate(results)]
            if all(status == "waiting_url" for status in results):
                source["status"] = "waiting_url"

    def image_text(self, source, image, frame, remote_status):
        """Supplement structured remote OCR with the native image text parser."""
        from langchain_community.document_loaders.blob_loaders import Blob
        if self.request_guard:
            self.request_guard()
        started = perf_counter()
        receipt = dict(frame=frame, method='langchain.RapidOCRBlobParser',
                       remote_status=remote_status)
        source['metadata'].setdefault('image_text_ocr', []).append(receipt)
        try:
            buffer = io.BytesIO()
            image.convert('RGB').save(buffer, format='PNG')
            with self.image_parser_lock:
                if self.image_parser is None:
                    from langchain_community.document_loaders.parsers.images import RapidOCRBlobParser
                    from rapidocr import RapidOCR
                    self.image_parser = RapidOCRBlobParser()
                    self.image_parser.ocr = RapidOCR(params={
                        'EngineConfig.onnxruntime.intra_op_num_threads': 2,
                        'EngineConfig.onnxruntime.inter_op_num_threads': 1})
                documents = self.image_parser.parse(Blob.from_data(buffer.getvalue(), mime_type='image/png'))
            text = '\n'.join(doc.page_content for doc in documents).strip()
            receipt['status'] = 'ready' if text else 'empty'
            if not text:
                self.issue(source, 'image_text_empty', '本地图片全文未识别到文字；保留远程OCR结果。', blocking=False)
                return
            path = self.output_dir / 'ocr' / source['source_id'] / f'frame-{frame}-local.txt'
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(text, encoding='utf-8')
            self.artifact(path, 'image_text_ocr', source, original_frame=frame)
            self.block(source, f'frame:{frame}/local-text', text, kind='image_ocr_text', frame=frame,
                extraction_method='langchain.RapidOCRBlobParser/rapidocr-3.10.0', priority=20,
                source_content_id=source['content_id'], ocr_status=remote_status,
                authoritative=remote_status == 'ready')
        except Exception as exc:
            receipt.update(status='failed', error=type(exc).__name__)
            self.issue(source, 'image_text_failed', f'本地图片全文补充失败：{type(exc).__name__}: {exc}', blocking=False)
        finally:
            receipt['seconds'] = round(perf_counter() - started, 6)

    def _ocr_response(self, source, path, source_url, mapping, receipt):
        import httpx
        cache = self.ocr_cache
        cache_key = _cache.OCRResponseCache.key(_digest(path.read_bytes()), self.ocr_endpoint, self.ocr_cache_version)
        receipt['cache_key'] = cache_key
        started = perf_counter()
        with cache.singleflight(cache_key, self.request_guard) if cache else nullcontext():
            receipt['cache_wait_seconds'] = round(perf_counter() - started, 6)
            cached = cache.read(cache_key) if cache else None
            if cached:
                wire, created_at = cached
                receipt.update(cache_hit=True, cached_at=created_at)
                self.progress(source, 'cache_hit', cache_key=cache_key)
                return wire, json.loads(wire)
            self.progress(source, 'uploading')
            started = perf_counter()
            url = self._ocr_url(source, path, source_url, mapping, receipt['anchor'])
            receipt['upload_and_sign_seconds'] = round(perf_counter() - started, 6)
            if not url:
                return None, None
            payload = {'file': url, 'task': 'orderChange', 'callBack': '', 'taskType': ''}
            if self.request_guard:
                self.request_guard()
            self.progress(source, 'requesting')
            started = perf_counter()
            try:
                response = (self.ocr_request(endpoint=self.ocr_endpoint, payload=payload,
                    timeout=httpx.Timeout(300.0, connect=10.0)) if self.ocr_request else
                    httpx.post(self.ocr_endpoint, json=payload, timeout=httpx.Timeout(300.0, connect=10.0)))
            finally:
                receipt['http_seconds'] = round(perf_counter() - started, 6)
            if isinstance(response, httpx.Response):
                wire = response.content
                receipt['http_status'] = response.status_code
                # Preserve HTTP failures too; never store them in the shared cache.
                raw = self.output_dir / 'ocr' / source['source_id'] / (_safe_name(receipt['anchor']) + '.response')
                raw.parent.mkdir(parents=True, exist_ok=True)
                raw.write_bytes(wire)
                self.artifact(raw, 'ocr_wire_response', source, **mapping_without_source(mapping))
                response.raise_for_status()
                result = response.json()
            else:
                result = response
                wire = json.dumps(result, ensure_ascii=False).encode('utf-8')
            documents = result.get('data', {}).get('list') if isinstance(result, dict) and isinstance(result.get('data'), dict) else None
            valid = (isinstance(result, str) and bool(result.strip())) or (
                isinstance(result, dict) and result.get('success') is True and isinstance(documents, list)
                and bool(documents) and all(isinstance(item, dict) and item for item in documents))
            if cache and valid:
                try:
                    receipt['cached_at'] = cache.write(cache_key, wire)
                except OSError as exc:
                    # A reusable cache is optional; the task still owns the raw receipt.
                    receipt['cache_error'] = f'{type(exc).__name__}: {exc}'
            return wire, result

    def _ocr_url(self, source, path, source_url, mapping, anchor):
        url = source_url
        if url is None and self.url_resolver:
            url = self.url_resolver(path, mapping)
        if not url and self.obs_options.get("enabled") is True:
            try:
                # SDK initialization and receipt slicing form one transaction;
                # slow OCR requests run concurrently outside this lock.
                with self.publisher_lock:
                    if self.publisher is None:
                        self.publisher = OBSPublisher(self.obs_options, self.request_guard)
                    start = len(self.publisher.receipts)
                    try:
                        url = self.publisher.publish(path, mapping)
                    finally:
                        source["metadata"].setdefault("obs_objects", []).extend(self.publisher.receipts[start:])
            except Exception as exc:
                self.issue(source, "obs_upload_failed", str(exc), scope=anchor)
                return None
        parsed = urlparse(url or "")
        if not url or parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.hostname.lower() in {"localhost", "127.0.0.1", "::1"}:
            self.issue(source, "ocr_waiting_url", "指定OCR服务需要可下载的HTTP(S)文件URL；本地原件尚未发布。")
            return None
        return url

    def ocr(self, source: dict, path: Path, source_url: str | None, mapping: dict, anchor: str) -> str:
        config_digest = _digest(json.dumps({"skill_version": VERSION, "endpoint": self.ocr_endpoint, "text_pointer": self.ocr_text_pointer, "response_format": self.ocr_response_format, "task": "orderChange"}, sort_keys=True).encode())
        source["metadata"]["ocr_config_digest"] = config_digest
        source["metadata"]["ocr_service_version"] = "unknown"
        import httpx
        raw_path = self.output_dir / "ocr" / source["source_id"] / (_safe_name(anchor) + ".json")
        raw_path.parent.mkdir(parents=True, exist_ok=True)
        started = perf_counter()
        receipt = dict(anchor=anchor, cache_hit=False, http_seconds=0.0,
            upload_and_sign_seconds=0.0, started_at=datetime.now(timezone.utc).isoformat(), **mapping_without_source(mapping))
        source['metadata'].setdefault('ocr_requests', []).append(receipt)
        try:
            wire, result = self._ocr_response(source, path, source_url, mapping, receipt)
            if wire is None:
                return 'failed' if any(i['code'] == 'obs_upload_failed' and i['source_id'] == source['source_id'] for i in self.issues) else 'waiting_url'
            raw_path.with_suffix(".response").write_bytes(wire)
            if not any(item['path'] == str(raw_path.with_suffix('.response').resolve()) for item in self.artifacts):
                self.artifact(raw_path.with_suffix(".response"), "ocr_wire_response", source, **mapping_without_source(mapping))
            raw_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
            self.artifact(raw_path, "ocr_raw", source, cache_key=receipt['cache_key'], ocr_config_digest=config_digest, service_version="unknown", **mapping_without_source(mapping))
            if self.ocr_response_format == "orderChange" and self.ocr_text_pointer is None:
                # Verified against the real service on 2026-10-21: it returns
                # structured documents, not raw OCR text. Preserve exact values.
                if not isinstance(result, dict) or result.get("success") is not True:
                    self.issue(source, "ocr_service_failed", "OCR未确认处理成功；原始响应已保存。", scope=anchor)
                    return "failed"
                documents = result.get("data", {}).get("list") if isinstance(result.get("data"), dict) else None
                if not isinstance(documents, list) or not documents or not all(isinstance(item, dict) and item for item in documents):
                    self.issue(source, "ocr_response_invalid", "orderChange响应缺少非空data.list单据对象；原始响应已保存。", scope=anchor)
                    return "partial"
                for index, document in enumerate(documents):
                    metadata = {"kind": "ocr_structured", "raw_response": str(raw_path.resolve()),
                        "response_pointer": f"/data/list/{index}", "original_pages": mapping.get("original_pages"),
                        "frame": mapping.get("frame"), "response_format": "orderChange",
                        "priority": 0 if source.get('format') == 'image' else 30}
                    if mapping.get("original_pages") and len(mapping["original_pages"]) == 1:
                        metadata["page"] = mapping["original_pages"][0]
                    self.block(source, f"{anchor}/document:{index + 1}",
                        json.dumps(document, ensure_ascii=False, indent=2), **metadata)
                return "ready"
            text = result
            if self.ocr_text_pointer is not None:
                for key in self.ocr_text_pointer.strip("/").split("/") if self.ocr_text_pointer else []:
                    key = key.replace("~1", "/").replace("~0", "~")
                    text = text[int(key)] if isinstance(text, list) else text[key]
            if not isinstance(text, str) or not text.strip():
                self.issue(source, "ocr_response_unmapped", "原响应已保存；需要经真实响应确认的文本JSON Pointer，不能猜测页/表/坐标字段。", scope=anchor)
                return "partial"
            metadata = {"kind": "ocr_text", "raw_response": str(raw_path.resolve()), "original_pages": mapping.get("original_pages"), "frame": mapping.get("frame"),
                        "priority": 0 if source.get('format') == 'image' else 30}
            if mapping.get("original_pages") and len(mapping["original_pages"]) == 1:
                metadata["page"] = mapping["original_pages"][0]
            self.block(source, anchor, text, **metadata)
            return "ready"
        except httpx.TimeoutException as exc:
            self.issue(source, "ocr_timeout", str(exc), scope=anchor)
        except httpx.HTTPStatusError as exc:
            self.issue(source, "ocr_http_error", f"HTTP {exc.response.status_code}", scope=anchor)
        except httpx.RequestError as exc:
            self.issue(source, "ocr_unreachable", str(exc), scope=anchor)
        except (KeyError, IndexError, TypeError, ValueError) as exc:
            self.issue(source, "ocr_response_invalid", f"{type(exc).__name__}: {exc}", scope=anchor)
        finally:
            receipt['elapsed_seconds'] = round(perf_counter() - started, 6)
            receipt['finished_at'] = datetime.now(timezone.utc).isoformat()
            self.progress(source, 'ocr_complete', receipt=deepcopy(receipt))
        return "failed"

    def office_converter(self):
        return (self.doc_converter or shutil.which("soffice.com") or
                shutil.which("soffice") or shutil.which("libreoffice"))

    def convert_office(self, source: dict, path: Path, extension: str, filter_name: str,
                       purpose: str) -> Path:
        """Delegate conversion/recalculation to LibreOffice and retain both byte versions."""
        converter = self.office_converter()
        if not converter:
            raise ValueError("需要配置LibreOffice soffice转换器")
        target_dir = self.output_dir / "office-conversion" / source["source_id"] / purpose
        target_dir.mkdir(parents=True, exist_ok=True)
        record = dict(purpose=purpose, original_path=str(path.resolve()),
                      original_sha256=_digest(path.read_bytes()),
                      original_content_id=source["content_id"],
                      converter=str(converter), format=extension)
        creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        try:
            if self.request_guard:
                self.request_guard()
            if self.converter_version is None:
                version = subprocess.run([str(converter), "--version"], capture_output=True,
                                         text=True, timeout=10, creationflags=creationflags)
                if version.returncode != 0:
                    raise ValueError("LibreOffice版本查询失败")
                self.converter_version = version.stdout.strip() or version.stderr.strip()
            record["converter_version"] = self.converter_version
            with tempfile.TemporaryDirectory(prefix="office-profile-") as profile:
                command = [str(converter), f"-env:UserInstallation={Path(profile).as_uri()}",
                           "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
                           "--convert-to", f"{extension}:{filter_name}", "--outdir",
                           str(target_dir.resolve()), str(path.resolve())]
                record["command"] = command
                result = subprocess.run(command, capture_output=True, text=True, timeout=120,
                                        creationflags=creationflags)
            record.update(returncode=result.returncode, stdout=result.stdout, stderr=result.stderr)
            converted = target_dir / (path.stem + "." + extension)
            if result.returncode != 0 or not converted.is_file():
                raise ValueError("LibreOffice未生成期望的派生文件")
            record.update(derived_path=str(converted.resolve()),
                          derived_sha256=_digest(converted.read_bytes()), status="ready")
            artifact = self.artifact(converted, "office_" + purpose, source,
                                     original_content_id=source["content_id"],
                                     original_sha256=record["original_sha256"],
                                     converter_version=self.converter_version, purpose=purpose)
            source["metadata"].setdefault("office_conversions", []).append(dict(artifact))
            return converted
        except Exception as exc:
            record.update(status="failed", error=f"{type(exc).__name__}: {exc}")
            if isinstance(exc, subprocess.TimeoutExpired):
                record["stdout"] = str(exc.stdout or "")
                record["stderr"] = str(exc.stderr or "")
            raise
        finally:
            log = target_dir / "conversion.json"
            log.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
            self.artifact(log, "conversion_log", source, purpose=purpose,
                          original_content_id=source["content_id"])

    def excel(self, source: dict, path: Path, kind: str):
        if kind == "xls":
            if self.office_converter():
                try:
                    converted = self.convert_office(source, path, "xlsx", "Calc MS Excel 2007 XML",
                                                    "xls-to-xlsx")
                except Exception as exc:
                    self.issue(source, "xls_conversion_failed", f"{type(exc).__name__}: {exc}")
                else:
                    self.xlsx(source, converted, original_format="xls", original_path=path)
                    return
            import xlrd
            from openpyxl.utils import get_column_letter
            book = xlrd.open_workbook(str(path), formatting_info=True)
            source["metadata"]["sheets"] = []
            for sheet in book.sheets():
                merged = [f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2)}{r2}"
                          for r1, r2, c1, c2 in sheet.merged_cells]
                source["metadata"]["sheets"].append({"name": sheet.name, "hidden": bool(sheet.visibility),
                    "state": {0: "visible", 1: "hidden", 2: "veryHidden"}.get(sheet.visibility, "unknown"),
                    "merged_cells": merged, "tables": []})
                for row in range(sheet.nrows):
                    cells = []
                    for column in range(sheet.ncols):
                        cell = sheet.cell(row, column)
                        if cell.ctype == xlrd.XL_CELL_EMPTY:
                            continue
                        value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat() if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                        xf = book.xf_list[cell.xf_index]
                        number_format = book.format_map.get(xf.format_key)
                        cells.append({"column": column + 1, "coordinate": f"{get_column_letter(column + 1)}{row + 1}",
                            "value": _json_value(value), "cached_value": _json_value(value), "formula": None,
                            "formula_origin": "unavailable-in-xlrd", "cell_type": cell.ctype,
                            "number_format": number_format.format_str if number_format else None,
                            "hidden_column": bool(sheet.colinfo_map.get(column) and sheet.colinfo_map[column].hidden)})
                    if cells:
                        self.block(source, f"sheet:{sheet.name}/row:{row + 1}", "\t".join(str(item["value"]) for item in cells), kind="table_row", sheet=sheet.name, row=row + 1, cells=cells, hidden=bool(sheet.rowinfo_map.get(row) and sheet.rowinfo_map[row].hidden))
            # xlrd exposes cached values, not reliable formula lineage.
            self.issue(source, "xls_formula_lineage", "XLS使用缓存值读取；旧格式公式表达式/重算未验证，需核对涉及公式的金额。", blocking=True)
            book.release_resources()
            return
        self.xlsx(source, path)

    def xlsx(self, source: dict, path: Path, *, original_format="xlsx", original_path=None):
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        formulas = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        source["metadata"]["sheets"] = []
        original_cached = cached
        cache_path = path
        try:
            formula_cells = [(sheet.title, cell.coordinate) for sheet in formulas
                             for row in sheet.iter_rows() for cell in row if cell.data_type == "f"]
            missing_cache = any(cached[sheet][coordinate].value is None for sheet, coordinate in formula_cells)
            calculation = formulas.calculation
            stale_cache = bool(formula_cells and calculation and (
                calculation.calcMode == "manual" or calculation.fullCalcOnLoad or calculation.forceFullCalc))
            source["metadata"]["spreadsheet"] = dict(original_format=original_format,
                formula_count=len(formula_cells), epoch=formulas.epoch.isoformat(),
                calculation=dict(calculation.to_tree().attrib) if calculation else {},
                original_content_id=source["content_id"])
            if original_format == "xlsx" and formula_cells and (missing_cache or stale_cache) and self.office_converter():
                try:
                    cache_path = self.convert_office(source, path, "xlsx", "Calc MS Excel 2007 XML",
                                                     "formula-recalculation")
                    cached = load_workbook(cache_path, data_only=True)
                except Exception as exc:
                    self.issue(source, "formula_recalculation_failed", f"{type(exc).__name__}: {exc}")
            elif stale_cache and not missing_cache and not self.office_converter():
                self.issue(source, "formula_recalculation_required", "工作簿要求重新计算；需LibreOffice生成可追溯缓存。")
            source["metadata"]["spreadsheet"]["cache_source"] = str(cache_path.resolve())
            if getattr(formulas, "_external_links", []):
                self.issue(source, "external_workbook_links", "工作簿引用外部数据；本地计算不能确认外部数据版本，需复核。")
            for sheet in formulas:
                merges = [str(item) for item in sheet.merged_cells.ranges]
                tables = []
                for name in sheet.tables:
                    table = sheet.tables[name]
                    tables.append(dict(name=name, ref=table.ref, header_rows=table.headerRowCount,
                                       total_rows=table.totalsRowCount or 0))
                source["metadata"]["sheets"].append({"name": sheet.title, "state": sheet.sheet_state,
                    "hidden": sheet.sheet_state != "visible", "merged_cells": merges, "tables": tables,
                    "max_row": sheet.max_row, "max_column": sheet.max_column})
                for row in sheet.iter_rows():
                    cells = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        is_formula = cell.data_type == "f"
                        formula = (cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)) if is_formula else None
                        value = cached[sheet.title][cell.coordinate].value if is_formula else cell.value
                        if is_formula and value is None:
                            self.issue(source, "formula_cache_missing", f"{sheet.title}!{cell.coordinate} 有公式但无计算缓存。", scope=f"sheet:{sheet.title}!{cell.coordinate}")
                        if is_formula and formula is None:
                            self.issue(source, "formula_expression_unavailable", f"{sheet.title}!{cell.coordinate}为特殊公式，未取得可回读表达式。",
                                       scope=f"sheet:{sheet.title}!{cell.coordinate}")
                        if (cached[sheet.title][cell.coordinate].data_type if is_formula else cell.data_type) == "e":
                            self.issue(source, "formula_calculation_error", f"{sheet.title}!{cell.coordinate}计算结果为{value}。",
                                       scope=f"sheet:{sheet.title}!{cell.coordinate}")
                        cells.append({"coordinate": cell.coordinate, "column": cell.column,
                            "value": _json_value(value), "cached_value": _json_value(value) if is_formula else None,
                            "original_cached_value": _json_value(original_cached[sheet.title][cell.coordinate].value) if is_formula else None,
                            "formula": formula, "formula_origin": "libreoffice-conversion" if original_format == "xls" and formula else "original" if formula else None,
                            "formula_range": getattr(cell.value, "ref", None) if is_formula else None,
                            "value_origin": "libreoffice-recalculation" if formula and (cache_path != path or original_format == "xls") else "original",
                            "cell_type": cell.data_type, "number_format": cell.number_format,
                            "hidden_column": bool(sheet.column_dimensions[cell.column_letter].hidden),
                            "merged_range": next((str(merge) for merge in sheet.merged_cells.ranges if cell.coordinate in merge), None)})
                    if cells:
                        memberships = []
                        for table in tables:
                            _, first, _, last = range_boundaries(table["ref"])
                            if first <= row[0].row <= last:
                                memberships.append(dict(name=table["name"], ref=table["ref"],
                                    role="header" if row[0].row < first + (table["header_rows"] or 0) else
                                         "total" if row[0].row > last - table["total_rows"] else "data"))
                        self.block(source, f"sheet:{sheet.title}/row:{row[0].row}", "\t".join(f"{item['coordinate']}={item['value'] if item['value'] is not None else '[未计算] ' + str(item['formula'])}" for item in cells), kind="table_row", sheet=sheet.title, row=row[0].row, cells=cells, hidden=bool(sheet.row_dimensions[row[0].row].hidden), sheet_state=sheet.sheet_state,
                            tables=memberships, original_content_id=source["content_id"],
                            cache_source=str(cache_path.resolve()) if any(item["formula"] for item in cells) else None)
        finally:
            formulas.close()
            cached.close()
            if original_cached is not cached:
                original_cached.close()

    def docx(self, source: dict, path: Path, depth: int):
        from docx import Document as WordDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from lxml import etree
        document = WordDocument(path)
        seen_cells: dict[int, str] = {}

        def walk(container, prefix, *, part="word/document.xml", context=None):
            context = context or {}
            for i, item in enumerate(container.iter_inner_content(), 1):
                anchor = f"{prefix}/{i}"
                if isinstance(item, Paragraph):
                    self.block(source, anchor, item.text, kind="paragraph", part=part,
                        paragraph=i, xml_path=item._p.getroottree().getpath(item._p),
                        style=item.style.name if item.style else None,
                        manual_page_break=bool(item._p.xpath('.//w:br[@w:type="page"]')),
                        rendered_page_breaks=len(item._p.xpath('.//w:lastRenderedPageBreak')),
                        page_break_before=item.paragraph_format.page_break_before,
                        keep_with_next=item.paragraph_format.keep_with_next,
                        keep_together=item.paragraph_format.keep_together, **context)
                elif isinstance(item, Table):
                    for row_number, row in enumerate(item.rows, 1):
                        cells = []
                        grid_column = row.grid_cols_before + 1
                        for column, cell in enumerate(row.cells, 1):
                            identity = id(cell._tc)
                            cell_anchor = f"{anchor}/row:{row_number}/cell:{column}"
                            if identity in seen_cells:
                                cells.append({"column": column, "anchor": seen_cells[identity],
                                              "grid_column": grid_column, "merged_proxy": True})
                                grid_column += 1
                                continue
                            seen_cells[identity] = cell_anchor
                            # Keep XML elements alive to prevent identity recycling.
                            held_cells.append(cell._tc)
                            properties = cell._tc.tcPr
                            merge = properties.vMerge if properties is not None else None
                            cells.append({"column": column, "anchor": cell_anchor, "text": cell.text,
                                "grid_column": grid_column, "grid_span": cell.grid_span,
                                "vertical_merge": merge.val if merge is not None else None,
                                "merged_proxy": False})
                            walk(cell, cell_anchor, part=part,
                                 context=dict(table=anchor, row=row_number, cell=column,
                                              cell_anchor=cell_anchor))
                            grid_column += 1
                        self.block(source, f"{anchor}/row:{row_number}", "\t".join(cell.get("text", "[merged]") for cell in cells), kind="table_row", table=anchor, row=row_number, cells=cells, structure_only=True,
                            part=part, xml_path=row._tr.getroottree().getpath(row._tr),
                            grid_cols_before=row.grid_cols_before, grid_cols_after=row.grid_cols_after,
                            repeated_header=bool(row._tr.xpath('./w:trPr/w:tblHeader')),
                            cannot_split=bool(row._tr.xpath('./w:trPr/w:cantSplit')))

        held_cells: list = []
        walk(document, "word/body")
        parts_seen = set()
        for i, section in enumerate(document.sections, 1):
            for label, area in [("header", section.header), ("footer", section.footer), ("first_header", section.first_page_header), ("first_footer", section.first_page_footer), ("even_header", section.even_page_header), ("even_footer", section.even_page_footer)]:
                if str(area.part.partname) in parts_seen:
                    continue
                parts_seen.add(str(area.part.partname))
                walk(area, f"word/section:{i}/{label}", part=str(area.part.partname).lstrip("/"))
        with zipfile.ZipFile(path) as archive:
            word_xml = [name for name in archive.namelist() if name.startswith("word/") and name.endswith(".xml")]
            features = {"textboxes": 0, "tracked_changes": 0, "drawings": 0, "image_occurrences": 0}
            for name in word_xml:
                root = etree.fromstring(archive.read(name))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                features["tracked_changes"] += len(root.xpath("//w:ins | //w:del", namespaces=ns))
                features["drawings"] += len(root.xpath("//w:drawing | //w:pict", namespaces=ns))
                rel_name = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels")
                relationships = {}
                if rel_name in archive.namelist():
                    relationships = {item.get("Id"): item for item in etree.fromstring(archive.read(rel_name))}
                image_refs = root.xpath('//*[local-name()="blip" or local-name()="imagedata"]')
                rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
                for n, ref in enumerate(image_refs, 1):
                    rel_id = ref.get(f"{{{rel_ns}}}embed") or ref.get(f"{{{rel_ns}}}id") or ref.get(f"{{{rel_ns}}}link")
                    relation = relationships.get(rel_id)
                    if relation is None:
                        self.issue(source, "word_image_relationship", "图片关系缺失，无法定位原始图。", scope=f"{name}/image:{n}")
                        continue
                    if relation.get("TargetMode") == "External":
                        self.issue(source, "word_external_image", "外链图片未下载；需要调用方确认可用材料。", scope=f"{name}/image:{n}")
                        continue
                    target = posixpath.normpath(posixpath.join(posixpath.dirname(name), relation.get("Target", "")))
                    target = target.lstrip("/")
                    image_name = posixpath.basename(target)
                    part_path = self.output_dir / "word-images" / source["source_id"] / f"{_safe_name(name)}-{n}-{_safe_name(image_name)}"
                    part_path.parent.mkdir(parents=True, exist_ok=True)
                    part_path.write_bytes(archive.read(target))
                    child = self.parse(part_path, parent_id=source["source_id"], occurrence=f"{name}/image:{n}", depth=depth + 1, original_name=image_name)
                    child["metadata"].update({"word_part": name, "relationship_id": rel_id, "package_member": target})
                    self.artifact(part_path, "word_embedded_image", child, word_part=name)
                    features["image_occurrences"] += 1
                    if child["status"] != "ready":
                        self.issue(source, "word_image_incomplete", "嵌入图片OCR未完成。", scope=child["source_id"])
                for n, box in enumerate(root.xpath("//w:txbxContent", namespaces=ns), 1):
                    features["textboxes"] += 1
                    text = "\n".join(box.xpath(".//w:t/text()", namespaces=ns))
                    self.block(source, f"{name}/textbox:{n}", text, kind="textbox", part=name)
                if re.match(r"word/(footnotes|endnotes)\.xml$", name):
                    for n, note in enumerate(root, 1):
                        text = "\n".join(note.xpath(".//w:t/text()", namespaces=ns))
                        if text:
                            self.block(source, f"{name}/note:{n}", text, kind="note", part=name)
            source["metadata"].update({"word_features": features, "physical_pages_known": False})
            if features["tracked_changes"]:
                self.issue(source, "word_tracked_changes", "含修订内容；当前读取常规正文，须确认最终/原始视图后复核金额。")
            if features["drawings"] > features["image_occurrences"]:
                self.issue(source, "word_embedded_visuals", "含未映射为位图的绘图；须人工核对图形中的业务信息。")
        if self.office_preview and self.office_converter():
            try:
                preview = self.convert_office(source, path, "pdf", "writer_pdf_Export", "word-preview")
                self.word_preview(source, preview)
            except Exception as exc:
                self.issue(source, "word_preview_failed", f"Word预览转换失败：{type(exc).__name__}: {exc}")

    def word_preview(self, source: dict, path: Path):
        """Bind optional preview locations to this rendering, without changing evidence anchors."""
        from pypdf import PdfReader
        pages = [re.sub(r"\s+", " ", page.extract_text() or "").strip()
                 for page in PdfReader(path).pages]
        starts, cursor = [], 0
        for text in pages:
            starts.append((cursor, cursor + len(text)))
            cursor += len(text) + 1
        text = " ".join(pages)
        preview_hash = _digest(path.read_bytes())
        source["metadata"]["office_preview"] = dict(path=str(path.resolve()), format="pdf",
            sha256=preview_hash, page_count=len(pages), converter_version=self.converter_version,
            original_content_id=source["content_id"], authoritative=False)
        for block in source["blocks"]:
            quote = re.sub(r"\s+", " ", block["text"]).strip()
            start = text.find(quote) if len(quote) >= 5 and not block.get("structure_only") else -1
            if start < 0 or text.find(quote, start + 1) >= 0:
                continue
            end = start + len(quote)
            matches = [number for number, (first, last) in enumerate(starts, 1)
                       if start < last and end > first]
            if matches:
                block["preview"] = dict(pages=matches, artifact_sha256=preview_hash,
                    mapping="unique-normalized-text-match", authoritative=False)

    def doc(self, source: dict, path: Path, depth: int):
        converter = self.office_converter()
        if not converter:
            self.issue(source, "doc_converter_missing", "旧DOC需要LibreOffice soffice转换器。配置doc_converter或将其加入PATH。")
            source["status"] = "failed"
            return
        converted = self.convert_office(source, path, "docx", "Office Open XML Text", "doc-to-docx")
        self.artifact(converted, "converted_docx", source, original_content_id=source["content_id"])
        child = self.parse(converted, parent_id=source["source_id"], occurrence="converted-docx", depth=depth + 1)
        child["metadata"]["original_content_id"] = source["content_id"]
        child["metadata"]["conversion_version"] = self.converter_version
        source["metadata"]["derived_source_id"] = child["source_id"]
        if child["status"] != "ready":
            self.issue(source, "doc_conversion_review", "派生DOCX有待复核内容。")


def mapping_without_source(mapping: dict) -> dict:
    return {key: value for key, value in mapping.items() if key not in {"source_id", "kind"}}


_PREVIEW_LOCK = RLock()


def render_office_preview(payload: dict) -> dict:
    """Reuse Office conversion for an on-demand display artifact, without extraction."""
    path, output = Path(payload['path']), Path(payload['output_dir'])
    converter = payload.get('doc_converter')
    identity = dict(original_sha256=_digest(path.read_bytes()), converter=converter,
                    converter_mtime=Path(converter).stat().st_mtime_ns if converter else None,
                    presentation_version='1.0.0')
    directory = output / _digest(json.dumps(identity, sort_keys=True).encode())
    manifest = directory / 'preview.json'
    with _PREVIEW_LOCK:
        if manifest.is_file():
            preview = json.loads(manifest.read_text(encoding='utf-8'))
            cached = Path(preview['path'])
            if cached.is_file() and _digest(cached.read_bytes()) == preview['sha256']:
                return preview
        parser = Intake(directory, ocr_endpoint=OCR_ENDPOINT, ocr_request=None,
            url_resolver=None, ocr_text_pointer=None, pdf_pages_per_slice=20, max_depth=0,
            max_files=1, max_total_bytes=100*1024*1024, doc_converter=converter)
        source = dict(source_id='display', content_id=identity['original_sha256'],
                      blocks=[], metadata={})
        filter_name = 'calc_pdf_Export' if path.suffix.lower() in ('.xlsx', '.xls') else 'writer_pdf_Export'
        converted = parser.convert_office(source, path, 'pdf', filter_name, 'source-preview')
        parser.word_preview(source, converted)
        preview = dict(source['metadata']['office_preview'], presentation_version='1.0.0')
        manifest.write_text(json.dumps(preview, ensure_ascii=False), encoding='utf-8')
        return preview


def intake_document(path: str | Path, *, source_url: str | None = None, output_dir: str | Path | None = None, ocr_endpoint: str = OCR_ENDPOINT, ocr_request: Callable | None = None, url_resolver: Callable | None = None, ocr_text_pointer: str | None = None, pdf_pages_per_slice: int = 20, max_depth: int = 8, max_files: int = 200, max_total_bytes: int = 100 * 1024 * 1024, doc_converter: str | None = None, occurrence_id: str = "root", obs_options: dict | None = None, request_guard: Callable | None = None, ocr_response_format: str = "text", office_preview: bool = True, pdf_text_mode: str = 'full', attachment_concurrency: int = 1, ocr_cache_dir=None, ocr_cache_ttl=86400, ocr_cache_version='orderChange-v1', on_progress=None, image_text_ocr: bool = False) -> dict:
    """Read one material and recursive attachments, preserving incomplete ranges.

    ``ocr_request`` is an optional controlled HTTP callable accepting endpoint,
    payload, timeout. ``url_resolver(Path, mapping)`` supplies verified OCR URLs;
    enabled ``obs_options`` supplies Huawei OBS publishing when no URL exists.
    Text pointer is an explicit JSON Pointer verified against the configured OCR
    service's real response. ``request_guard`` runs before upload and OCR dispatch.
    """
    path = Path(path)
    output = Path(output_dir) if output_dir else path.parent / ".intake" / _digest(str(path.resolve()).encode())[:16]
    parser = Intake(output, ocr_endpoint=ocr_endpoint, ocr_request=ocr_request, url_resolver=url_resolver, ocr_text_pointer=ocr_text_pointer, pdf_pages_per_slice=pdf_pages_per_slice, max_depth=max_depth, max_files=max_files, max_total_bytes=max_total_bytes, doc_converter=doc_converter, obs_options=obs_options, request_guard=request_guard, ocr_response_format=ocr_response_format, office_preview=office_preview, pdf_text_mode=pdf_text_mode, attachment_concurrency=attachment_concurrency, ocr_cache_dir=ocr_cache_dir, ocr_cache_ttl=ocr_cache_ttl, ocr_cache_version=ocr_cache_version, on_progress=on_progress, image_text_ocr=image_text_ocr)
    try:
        parser.parse(path, source_url=source_url, occurrence=occurrence_id)
    finally:
        if parser.publisher:
            parser.publisher.close()
    statuses = [source["status"] for source in parser.sources]
    status = "ready" if all(item == "ready" for item in statuses) else "failed" if all(item == "failed" for item in statuses) else "waiting_url" if all(item in {"ready", "waiting_url"} for item in statuses) and "waiting_url" in statuses else "partial"
    coverage = [{"source_id": source["source_id"], "status": source["status"], "block_count": len(source["blocks"]), "ocr_ranges": source["metadata"].get("ocr_ranges", source["metadata"].get("ocr_frames", []))} for source in parser.sources]
    return {"skill": "document-intake", "version": VERSION, "status": status, "sources": parser.sources, "issues": parser.issues, "coverage": coverage, "artifacts": parser.artifacts}


class ShippingDocumentLoader(BaseLoader):
    """LangChain loader projection over the same intake implementation."""

    def __init__(self, path: str | Path, **options):
        self.path, self.options = path, options
        self.result: dict | None = None

    def lazy_load(self) -> Iterator[Document]:
        self.result = intake_document(self.path, **self.options)
        for source in self.result["sources"]:
            for block in source["blocks"]:
                yield Document(page_content=block["text"], metadata={"source": source["path"], "content_id": source["content_id"], "status": source["status"], **{key: value for key, value in block.items() if key != "text"}})


def main() -> int:
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("path", nargs="?")
    cli.add_argument("--input", help="JSON file containing path and intake options")
    cli.add_argument("--output", help="Result JSON; stdout is JSON when omitted")
    cli.add_argument("--output-dir")
    cli.add_argument("--source-url")
    cli.add_argument("--ocr-text-pointer")
    cli.add_argument("--ocr-response-format", choices=("text", "orderChange"))
    args = cli.parse_args()
    options = json.loads(Path(args.input).read_text(encoding="utf-8-sig")) if args.input else {"path": args.path}
    for key in ["output_dir", "source_url", "ocr_text_pointer", "ocr_response_format"]:
        if getattr(args, key) is not None:
            options[key] = getattr(args, key)
    if not options.get("path"):
        cli.error("Provide path or --input JSON containing path")
    result = intake_document(**options)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(text, encoding="utf-8")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(text)
    return 0 if result["status"] == "ready" else 2


if __name__ == "__main__":
    raise SystemExit(main())
