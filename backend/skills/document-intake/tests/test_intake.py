"""Offline contract checks. Generated fixtures are synthetic, not exam answers."""

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from email.message import EmailMessage
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter

spec = importlib.util.spec_from_file_location("document_intake", Path(__file__).parents[1] / "scripts" / "skill.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
intake_document = module.intake_document


def test_eml_ocr_attachments_overlap_with_shared_limit_and_preserve_failures(tmp_path):
    import io
    from threading import Barrier, Lock
    from concurrent.futures import ThreadPoolExecutor
    mail = EmailMessage()
    mail.set_content('Synthetic invoice USD 50')
    for i in range(3):
        buf = io.BytesIO()
        Image.new('RGB', (10, 10), 'white').save(buf, format='PNG')
        mail.add_attachment(buf.getvalue(), maintype='image', subtype='png', filename=f'{i}.png')
    path = tmp_path / 'bill.eml'
    path.write_bytes(mail.as_bytes())
    gate, lock, seen = Barrier(3, timeout=3), Lock(), []
    def ocr(**kwargs):
        with lock:
            seen.append(kwargs['payload']['file'])
        gate.wait()
        return httpx.Response(200, json='Fee USD 50', request=httpx.Request('POST', 'https://ocr.test'))
    result = intake_document(path, output_dir=tmp_path/'parsed', attachment_concurrency=3,
        url_resolver=lambda p, m: 'https://files.test/'+p.name, ocr_request=ocr)
    assert result['status'] == 'ready'
    assert len(seen) == 3 and len(result['sources']) == 4
    assert all(s['parent_source_id'] == result['sources'][0]['source_id'] for s in result['sources'][1:])
    assert len([a for a in result['artifacts'] if a['kind'] == 'ocr_raw']) == 3


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def image(self, name="image.png"):
        path = self.root / name
        Image.new("RGB", (30, 20), "white").save(path)
        return path

    def intake(self, path, **kwargs):
        return intake_document(path, output_dir=self.root / "artifacts", **kwargs)

    def test_image_waiting_url_is_not_success(self):
        result = self.intake(self.image())
        self.assertEqual(result["status"], "waiting_url")
        self.assertEqual(result["issues"][0]["code"], "ocr_waiting_url")
        self.assertEqual(result["sources"][0]["blocks"], [])

    def test_searchable_pdf_full_text_survives_reduced_ocr_after_real_receipt(self):
        path = Path(__file__).parents[3] / 'data/samples/PDF 示例账单2.pdf'
        calls = []
        def ocr(**kwargs):
            calls.append(kwargs)
            return {'success': True, 'data': {'list': [{'chargeList': []}]}}
        result = self.intake(path, source_url='https://files.example/bill.pdf',
                             ocr_response_format='orderChange', ocr_request=ocr)
        blocks = result['sources'][0]['blocks']
        full = [b for b in blocks if b['kind'] == 'pdf_embedded_text']
        self.assertEqual(len(calls), 1)
        self.assertEqual(len(full), 1)
        self.assertTrue(full[0]['authoritative'])
        self.assertEqual(full[0]['page'], 1)
        self.assertEqual(full[0]['ocr_status'], 'ready')
        self.assertEqual(full[0]['source_content_id'], result['sources'][0]['content_id'])
        for value in ('DC-108478', 'USD', 'COD', '1,410.40', '10-24-2026'):
            self.assertIn(value, full[0]['text'])
        self.assertTrue(any(b['kind'] == 'ocr_structured' for b in blocks))

    def test_pdf_text_cannot_hide_failed_or_missing_required_ocr(self):
        path = Path(__file__).parents[3] / 'data/samples/PDF 示例账单2.pdf'
        for options in ({}, dict(source_url='https://files.example/bill.pdf',
                ocr_response_format='orderChange', ocr_request=lambda **_: {'success': False})):
            result = self.intake(path, **options)
            self.assertNotEqual(result['status'], 'ready')
            self.assertFalse(any(b['authoritative'] and b['kind'] == 'pdf_embedded_text'
                                 for b in result['sources'][0]['blocks']))

    def test_ocr_only_leaves_pdf_supplement_choice_to_agent(self):
        path = Path(__file__).parents[3] / 'data/samples/PDF 示例账单2.pdf'
        result = self.intake(path, source_url='https://files.example/bill.pdf',
            pdf_text_mode='ocr_only', ocr_response_format='orderChange',
            ocr_request=lambda **_: {'success': True, 'data': {'list': [{'chargeList': []}]}})
        source = result['sources'][0]
        self.assertEqual(source['status'], 'ready')
        self.assertTrue(any(b['kind'] == 'ocr_structured' for b in source['blocks']))
        self.assertFalse(any(b['kind'] in ('pdf_embedded_text', 'pdf_layout_text') for b in source['blocks']))
        self.assertTrue(any('DC-108478' in b['text'] for b in source['blocks'] if b['kind'] == 'native_preview'))

    def test_hbl_embedded_text_keeps_fields_omitted_by_layout_mode(self):
        from email import policy
        from email.parser import BytesParser
        original = next((Path(__file__).parents[3] / 'data/samples').glob('*.eml'))
        mail = BytesParser(policy=policy.default).parsebytes(original.read_bytes())
        part = next(p for p in mail.walk() if (p.get_filename() or '').endswith('HBL.pdf'))
        path = self.root / 'hbl.pdf'
        path.write_bytes(part.get_payload(decode=True))
        result = self.intake(path, source_url='https://files.example/hbl.pdf',
            ocr_response_format='orderChange',
            ocr_request=lambda **_: {'success': True, 'data': {'list': [{'chargeList': []}]}})
        text = '\n'.join(b['text'] for b in result['sources'][0]['blocks'] if b['kind']=='pdf_embedded_text')
        for value in ('SSZSE26093128', 'OOLU4055331281', 'FREIGHT PREPAID', 'SHENZHEN XUELANG', 'TIIU8975975'):
            self.assertIn(value, text)

    def test_original_mbl_has_a_readable_header_layout_after_ocr(self):
        from email import policy
        from email.parser import BytesParser
        original = next((Path(__file__).parents[3] / 'data/samples').glob('*.eml'))
        mail = BytesParser(policy=policy.default).parsebytes(original.read_bytes())
        part = next(p for p in mail.walk() if (p.get_filename() or '').endswith('MBL.pdf'))
        path = self.root / 'mbl.pdf'
        path.write_bytes(part.get_payload(decode=True))
        result = self.intake(path, source_url='https://files.example/mbl.pdf',
            ocr_response_format='orderChange',
            ocr_request=lambda **_: {'success': True, 'data': {'list': [{'chargeList': []}]}})
        source = result['sources'][0]
        layouts = [b for b in source['blocks'] if b['kind'] == 'pdf_layout_text']
        self.assertEqual(len(layouts), 1)
        layout = layouts[0]
        self.assertTrue(layout['authoritative'])
        self.assertEqual(layout['page'], 1)
        self.assertEqual(layout['source_content_id'], source['content_id'])
        lines = layout['text'].splitlines()
        label = next(i for i, line in enumerate(lines) if 'BOOKING NUMBER' in line and 'B/L NO:' in line)
        self.assertIn('4055331281', lines[label + 1])
        self.assertIn('HZMHK331281', lines[label + 1])
        missing = self.intake(path)
        self.assertFalse(any(b['kind'] == 'pdf_layout_text' and b['authoritative']
                             for b in missing['sources'][0]['blocks']))

    def test_multiframe_tiff_each_frame_has_mapping(self):
        path = self.root / "frames.tiff"
        Image.new("RGB", (30, 20), "white").save(path, save_all=True, append_images=[Image.new("RGB", (20, 30), "black")])
        mappings = []

        def resolve(part, mapping):
            mappings.append(mapping)
            return "https://files.example/" + part.name

        result = self.intake(path, url_resolver=resolve, ocr_request=lambda **_: "synthetic frame text")
        self.assertEqual(result["status"], "ready")
        self.assertEqual([mapping["frame"] for mapping in mappings], [1, 2])
        self.assertEqual([block["frame"] for block in result["sources"][0]["blocks"]], [1, 2])

    def test_ocr_exact_request_and_raw_evidence(self):
        calls = []

        def request(**kwargs):
            calls.append(kwargs)
            return {"data": {"text": "Handling Fee USD 10"}}

        result = self.intake(self.image(), source_url="https://files.example/material.png", ocr_request=request, ocr_text_pointer="/data/text")
        self.assertEqual(result["status"], "ready")
        self.assertEqual(calls[0]["payload"], {"file": "https://files.example/material.png", "task": "orderChange", "callBack": "", "taskType": ""})
        self.assertEqual(calls[0]["timeout"].connect, 10)
        self.assertEqual(calls[0]["timeout"].read, 300)
        block = result["sources"][0]["blocks"][0]
        self.assertNotIn("page", block)
        self.assertEqual(block["text"], "Handling Fee USD 10")
        self.assertEqual(json.loads(Path(block["raw_response"]).read_text(encoding="utf-8")), {"data": {"text": "Handling Fee USD 10"}})

    def test_ocr_unknown_schema_and_http_error_keep_raw(self):
        path = self.image()
        result = self.intake(path, source_url="https://files.example/a.png", ocr_request=lambda **_: {"unverified": "10"})
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["issues"][0]["code"], "ocr_response_unmapped")
        result = self.intake(path, source_url="https://files.example/a.png", ocr_request=lambda **_: httpx.Response(503, content=b"OCR unavailable", request=httpx.Request("POST", module.OCR_ENDPOINT)))
        self.assertEqual(result["issues"][0]["code"], "ocr_http_error")
        self.assertTrue(any(item["kind"] == "ocr_wire_response" for item in result["artifacts"]))

    def test_pdf_slices_keep_original_page_mapping(self):
        path = self.root / "long.pdf"
        writer = PdfWriter()
        for _ in range(5):
            writer.add_blank_page(width=200, height=300)
        with path.open("wb") as stream:
            writer.write(stream)
        mappings = []

        def resolve(part, mapping):
            mappings.append(mapping)
            self.assertEqual(len(PdfReader(part).pages), len(mapping["original_pages"]))
            return f"https://files.example/{part.name}"

        result = self.intake(path, pdf_pages_per_slice=2, url_resolver=resolve, ocr_request=lambda **_: "synthetic OCR")
        self.assertEqual(result["status"], "ready")
        self.assertEqual([mapping["original_pages"] for mapping in mappings], [[1, 2], [3, 4], [5]])
        blocks = result["sources"][0]["blocks"]
        self.assertEqual(len([block for block in blocks if block["kind"] == "native_preview"]), 5)
        self.assertTrue(all(not block["authoritative"] for block in blocks if block["kind"] == "native_preview"))
        ocr = [block for block in blocks if block["kind"] == "ocr_text"]
        self.assertNotIn("page", ocr[0])
        self.assertEqual(ocr[2]["page"], 5)

    def test_mail_duplicate_content_keeps_two_occurrences_and_nested_mail(self):
        data = self.image().read_bytes()
        message = EmailMessage()
        message["Subject"] = "Synthetic fee notice"
        message["From"] = "a@example.test"
        message["To"] = "b@example.test"
        message.set_content("Current amount USD 10")
        message.add_alternative("<html><body>Current amount USD 10</body></html>", subtype="html")
        message.add_attachment(data, maintype="image", subtype="png", filename="same.png")
        message.add_attachment(data, maintype="image", subtype="png", filename="same.png")
        nested = EmailMessage()
        nested.set_content("Forwarded old amount USD 9")
        message.add_attachment(nested)
        path = self.root / "sample.eml"
        path.write_bytes(message.as_bytes())
        result = self.intake(path)
        images = [source for source in result["sources"] if source["format"] == "image"]
        self.assertEqual(len(images), 2)
        self.assertEqual(images[0]["content_id"], images[1]["content_id"])
        self.assertNotEqual(images[0]["source_id"], images[1]["source_id"])
        self.assertNotEqual(images[0]["path"], images[1]["path"])
        root = result["sources"][0]
        self.assertEqual(len([block for block in root["blocks"] if block["kind"] == "mail_body"]), 1)
        self.assertEqual(root["blocks"][0]["anchor"], "mail/headers")
        self.assertEqual(root["metadata"]["headers"]["Subject"], "Synthetic fee notice")
        self.assertTrue(any(source["format"] == "eml" and source["parent_source_id"] == root["source_id"] for source in result["sources"][1:]))
        self.assertEqual(root["status"], "partial")

    def test_excel_formula_cache_and_hidden_merge_metadata(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "费用"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "USD"
        sheet["A2"] = "Handling Fee"
        sheet["B2"] = 10
        sheet["B3"] = "=SUM(B2:B2)"
        sheet.row_dimensions[2].hidden = True
        path = self.root / "fee.xlsx"
        book.save(path)
        result = self.intake(path)
        self.assertEqual(result["status"], "partial")
        self.assertTrue(any(issue["code"] == "formula_cache_missing" for issue in result["issues"]))
        source = result["sources"][0]
        self.assertEqual(source["metadata"]["sheets"][0]["merged_cells"], ["A1:B1"])
        self.assertTrue(source["blocks"][1]["hidden"])
        self.assertIsNone(source["blocks"][2]["cells"][0]["value"])

    def test_word_order_nested_table_and_merged_cell_identity(self):
        document = Document()
        document.add_paragraph("Start")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Merged fee USD 10"
        table.cell(0, 0).merge(table.cell(0, 1))
        nested = table.cell(1, 0).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Nested amount 5"
        document.add_paragraph("End")
        path = self.root / "fee.docx"
        document.save(path)
        result = self.intake(path)
        self.assertEqual(result["status"], "ready")
        blocks = result["sources"][0]["blocks"]
        content = [block for block in blocks if not block.get("structure_only")]
        self.assertEqual(sum("Merged fee" in block["text"] for block in content), 1)
        self.assertTrue(any("Nested amount 5" in block["text"] for block in content))
        self.assertEqual(content[0]["text"], "Start")
        self.assertIn("End", "\n".join(block["text"] for block in content))
        self.assertFalse(any("page" in block for block in blocks))

    def test_word_embedded_image_occurrences_preserved(self):
        image = self.image()
        document = Document()
        document.add_picture(str(image))
        document.add_picture(str(image))
        path = self.root / "embedded.docx"
        document.save(path)
        result = self.intake(path)
        self.assertEqual(result["status"], "partial")
        images = [source for source in result["sources"] if source["format"] == "image"]
        self.assertEqual(len(images), 2)
        self.assertEqual(images[0]["content_id"], images[1]["content_id"])
        self.assertNotEqual(images[0]["source_id"], images[1]["source_id"])
        self.assertTrue(all(source["status"] == "waiting_url" for source in images))

    def test_legacy_doc_missing_converter_is_explicit(self):
        path = self.root / "legacy.doc"
        path.write_bytes(b"\xd0\xcf\x11\xe0" + b"synthetic placeholder")
        with patch.object(module.shutil, "which", return_value=None):
            result = self.intake(path)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["issues"][0]["code"], "doc_converter_missing")

    def test_corrupt_attachment_preserves_sibling(self):
        message = EmailMessage()
        message.set_content("Body")
        message.add_attachment(b"%PDF-invalid", maintype="application", subtype="pdf", filename="broken.pdf")
        message.add_attachment(b"Fee USD 12", maintype="text", subtype="plain", filename="fee.txt")
        path = self.root / "sample.eml"
        path.write_bytes(message.as_bytes())
        result = self.intake(path)
        self.assertEqual(result["status"], "partial")
        self.assertTrue(any(source["status"] == "failed" for source in result["sources"]))
        self.assertTrue(any(source["text"] == "Fee USD 12" for source in result["sources"]))

    def test_loader_uses_official_document_projection(self):
        path = self.root / "fee.txt"
        path.write_text("Fee USD 10", encoding="utf-8")
        loader = module.ShippingDocumentLoader(path, output_dir=self.root / "out")
        docs = loader.load()
        self.assertEqual(docs[0].page_content, "Fee USD 10")
        self.assertEqual(docs[0].metadata["source_id"], loader.result["sources"][0]["source_id"])

    def test_cli_runs_from_outside_project_and_json_only_stdout(self):
        path = self.root / "fee.txt"
        path.write_text("Fee USD 10", encoding="utf-8")
        script = Path(module.__file__).resolve()
        response = subprocess.run([sys.executable, str(script), str(path), "--output-dir", str(self.root / "cli-output")], cwd=self.root, capture_output=True, encoding="utf-8", check=True)
        self.assertEqual(json.loads(response.stdout)["status"], "ready")
        self.assertEqual(response.stderr, "")

    def test_limits_make_unhandled_occurrence_visible(self):
        message = EmailMessage()
        message.set_content("Body")
        message.add_attachment(b"Fee USD 12", maintype="text", subtype="plain", filename="fee.txt")
        path = self.root / "sample.eml"
        path.write_bytes(message.as_bytes())
        result = self.intake(path, max_files=1)
        self.assertTrue(any(issue["code"] == "intake_limit" for issue in result["issues"]))
        self.assertEqual(len(result["sources"]), 2)


if __name__ == "__main__":
    unittest.main()
