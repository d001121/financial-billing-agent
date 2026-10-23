"""Targeted Markdown-first checks; all OCR responses are synthetic."""
import importlib.util
import json
import pytest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image

spec = importlib.util.spec_from_file_location('intake_markdown_tests', Path(__file__).parents[1]/'scripts/skill.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def check_mapping(source):
    md = source['markdown']
    assert Path(md['path']).read_text(encoding='utf-8') == md['text']
    blocks = {b['anchor']: b for b in source['blocks']}
    for span in md['spans']:
        raw = blocks[span['anchor']]['text'][span['start']:span['end']]
        assert raw == span['quote']
        assert 1 <= span['line'] <= md['line_count']
    return md


def test_all_supported_text_office_and_image_inputs_produce_markdown(tmp_path):
    word = Document()
    word.add_paragraph('Invoice overseas')
    table = word.add_table(rows=2, cols=2)
    for cell, value in zip([c for r in table.rows for c in r.cells], ['Fee', 'Amount', '运费', 'USD 50']):
        cell.text = value
    word.save(tmp_path/'bill.docx')
    book = Workbook()
    book.active.append(['Fee', 'Currency', 'Amount'])
    book.active.append(['Freight', 'EUR', 125])
    book.save(tmp_path/'bill.xlsx')
    (tmp_path/'bill.txt').write_text('  Freight  USD 50  \r\n\r\n', encoding='utf-8')
    Image.new('RGB', (5, 5)).save(tmp_path/'bill.png')
    for name in ('bill.docx', 'bill.xlsx', 'bill.txt', 'bill.png'):
        result = module.intake_document(tmp_path/name, output_dir=tmp_path/(name+'.out'),
            office_preview=False, source_url='https://files.test/bill',
            ocr_request=lambda **_: 'Freight USD 50')
        source = result['sources'][0]
        md = check_mapping(source)
        assert md['text'].strip() and md['spans']
        assert any(a['kind'] == 'markdown' for a in result['artifacts'])
        assert any(a['kind'] == 'markdown_map' for a in result['artifacts'])


def test_structured_ocr_becomes_readable_table_without_claiming_verbatim_fulltext(tmp_path):
    Image.new('RGB', (5, 5)).save(tmp_path/'bill.png')
    result = module.intake_document(tmp_path/'bill.png', output_dir=tmp_path/'out',
        source_url='https://files.test/bill', ocr_response_format='orderChange',
        ocr_request=lambda **_: dict(success=True, data=dict(list=[dict(
            invoiceNo='INV-9', chargeList=[dict(name='Freight', amount='1.234,56', currency='EUR')])])))
    md = check_mapping(result['sources'][0])
    assert '| Field | Value |' in md['text']
    assert 'chargeList/0/amount' in md['text'] and '1.234,56' in md['text']
    assert md['representations'][0]['kind'] == 'ocr_structured'
    assert md['representations'][0]['verbatim_fulltext'] is False


def test_failed_ocr_still_produces_explicit_incomplete_markdown(tmp_path):
    Image.new('RGB', (5, 5)).save(tmp_path/'bill.png')
    result = module.intake_document(tmp_path/'bill.png', output_dir=tmp_path/'out')
    md = check_mapping(result['sources'][0])
    assert md['status'] != 'ready'
    assert not md['spans']
    assert 'waiting_url' in md['text']


def test_pdf_markdown_includes_original_page_text_and_service_receipt(tmp_path):
    path = Path(__file__).parents[3]/'data/samples/PDF 示例账单2.pdf'
    calls = []
    def ocr(**kwargs):
        calls.append(kwargs)
        return dict(success=True, data=dict(list=[dict(chargeList=[])]))
    result = module.intake_document(path, output_dir=tmp_path/'out',
        source_url='https://files.test/bill', ocr_response_format='orderChange',
        ocr_request=ocr, pdf_text_mode='markdown')
    md = check_mapping(result['sources'][0])
    assert calls and 'TOTAL AMOUNT 1,410.40' in md['text']
    assert any(r.get('page') == 1 for r in md['representations'])
    assert any(r['kind'] == 'pdf_layout_text' for r in md['representations'])
    assert any(r['kind'] == 'ocr_structured' for r in md['representations'])


def test_pdf_layout_financial_rows_are_primary_and_keep_three_lines(tmp_path):
    path = Path(__file__).parents[3]/'data/samples/PDF 示例账单2.pdf'
    def ocr(**kwargs):
        return dict(success=True, data=dict(list=[dict(chargeList=[
            dict(chargeItem='HANDLING CHARGE', amount='100.00'),
            dict(chargeItem="4 ADD'L DETENTION FREE TIME CHARGE", amount='1310.40')])]))
    result = module.intake_document(path, output_dir=tmp_path/'out',
        source_url='https://files.test/bill', ocr_response_format='orderChange',
        ocr_request=ocr, pdf_text_mode='markdown')
    source = result['sources'][0]
    financial = [b for b in source['blocks'] if b['kind'] == 'pdf_layout_financial']
    assert len(financial) == 1
    assert [row['amount'] for row in financial[0]['rows']] == ['65.00', '1310.40', '35.00']
    md = check_mapping(source)
    kinds = [item['kind'] for item in md['representations']]
    assert kinds.index('pdf_layout_financial') < kinds.index('pdf_layout_text')
    assert kinds.index('pdf_layout_text') < kinds.index('pdf_embedded_text')
    assert kinds.index('pdf_embedded_text') < kinds.index('ocr_structured')
    assert '1,310.40' in md['text'] and '35.00' in md['text']


def test_pdf_recheck_sends_only_selected_original_page_to_ocr_and_reuses_receipt(tmp_path):
    from hashlib import sha256
    from pypdf import PdfReader, PdfWriter
    path = tmp_path/'two.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.add_blank_page(width=400, height=500)
    writer.write(path)
    source = dict(source_id='pdf', format='pdf', status='partial', metadata={}, blocks=[],
        content_id='sha256:' + sha256(path.read_bytes()).hexdigest())
    sent = []
    def resolver(selected, mapping):
        page = PdfReader(selected)
        assert len(page.pages) == 1 and page.pages[0].mediabox.width == 400
        assert mapping['original_pages'] == [2]
        sent.append(mapping)
        return 'https://files.test/selected.pdf'
    options = dict(path=path, output_dir=tmp_path/'out', reason='原页字段缺口', pages=[2],
        url_resolver=resolver, ocr_request=lambda **_: 'Amount EUR 50')
    result = module.supplement_ocr(source, **options)
    assert len(sent) == 1 and result['reads'][0]['method'] == 'specified_ocr'
    block = result['source']['blocks'][0]
    assert block['page'] == 2 and block['original_pages'] == [2]
    assert block['text'] == 'Amount EUR 50' and block['kind'] == 'ocr_text'
    assert Path(block['raw_response']).is_file()
    assert 'Amount EUR 50' in check_mapping(result['source'])['text']
    assert result['source']['status'] == 'partial'  # other failed ranges remain unresolved
    assert source['blocks'] == []
    repeated = module.supplement_ocr(result['source'], **dict(options, output_dir=tmp_path/'again'))
    assert repeated['added_anchors'] == [] and repeated['reads'][0]['reused']
    assert len(sent) == 1 and not any('base64' in b for b in result['source']['blocks'])


@pytest.mark.parametrize('pages', [[0], [3], [True], [], None])
def test_pdf_recheck_rejects_invalid_page_before_ocr(tmp_path, pages):
    from hashlib import sha256
    from pypdf import PdfWriter
    path = tmp_path/'one.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.write(path)
    source = dict(source_id='pdf', format='pdf', metadata={}, blocks=[],
        content_id='sha256:' + sha256(path.read_bytes()).hexdigest())
    with pytest.raises(ValueError, match='真实页'):
        module.supplement_ocr(source, path=path, output_dir=tmp_path/'out', pages=pages,
            reason='check', ocr_request=lambda **_: pytest.fail('must not dispatch OCR'))


def test_image_frame_recheck_uses_ocr_and_preserves_failure(tmp_path):
    from hashlib import sha256
    from io import BytesIO
    import httpx
    path = tmp_path/'frames.tiff'
    Image.new('RGB', (16, 20), 'red').save(path, save_all=True,
        append_images=[Image.new('RGB', (16, 20), 'blue')])
    source = dict(source_id='image', format='image', status='ready', metadata={}, blocks=[],
        content_id='sha256:' + sha256(path.read_bytes()).hexdigest())
    def resolver(selected, mapping):
        assert mapping['frame'] == 2
        assert Image.open(BytesIO(selected.read_bytes())).getpixel((0, 0)) == (0, 0, 255)
        return 'https://files.test/frame2.png'
    def failed(**_):
        raise httpx.ReadTimeout('synthetic timeout')
    options = dict(path=path, frame=2, reason='check', url_resolver=resolver)
    result = module.supplement_ocr(source, output_dir=tmp_path/'out', ocr_request=failed, **options)
    assert result['reads'][0]['status'] == 'failed' and result['source']['status'] == 'partial'
    assert result['added_anchors'] == [] and result['issues'][0]['code'] == 'ocr_timeout'
    assert check_mapping(result['source'])['status'] == 'partial'
    path.write_bytes(path.read_bytes() + b'changed')
    with pytest.raises(ValueError, match='内容已变化'):
        module.supplement_ocr(source, output_dir=tmp_path/'changed', **options)
