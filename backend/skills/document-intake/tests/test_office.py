"""Format-level Office regressions; no model or OCR requests."""

import hashlib
import importlib.util
import json
import os
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import pytest


ROOT = Path(__file__).parents[3]
FIXTURES = ROOT / "tests" / "fixtures" / "office"
spec = importlib.util.spec_from_file_location("office_intake", Path(__file__).parents[1] / "scripts" / "skill.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
CONVERTER = os.environ.get("OFFICE_TEST_CONVERTER") or str(ROOT / "vendor" / "libreoffice" / "program" / "soffice.com")
REAL_OFFICE = Path(CONVERTER).is_file()


def test_multisheet_tables_dates_and_hidden_cell_evidence(tmp_path):
    result = module.intake_document(FIXTURES / "fees-multisheet.xlsx", output_dir=tmp_path,
                                    office_preview=False)
    source = result["sources"][0]
    sheets = {sheet["name"]: sheet for sheet in source["metadata"]["sheets"]}
    assert sheets["Charges"]["tables"][0]["ref"] == "A4:E6"
    assert sheets["Additional charges"]["state"] == "hidden"
    assert sheets["Metadata"]["state"] == "veryHidden"
    blocks = {block["anchor"]: block for block in source["blocks"]}
    assert blocks["sheet:Charges/row:2"]["cells"][0]["value"] == "2026-10-21T00:00:00"
    assert blocks["sheet:Charges/row:6"]["hidden"] is True
    assert next(cell for cell in blocks["sheet:Charges/row:5"]["cells"] if cell["coordinate"] == "C5")["hidden_column"] is True
    assert next(cell for cell in blocks["sheet:Charges/row:5"]["cells"] if cell["coordinate"] == "D5")["value"] is None
    assert any(issue["code"] == "formula_cache_missing" for issue in result["issues"])
    assert blocks["sheet:Charges/row:5"]["tables"][0]["role"] == "data"
    json.dumps(result)


def test_word_original_anchors_and_cross_page_structure(tmp_path):
    result = module.intake_document(FIXTURES / "fees-cross-page.docx", output_dir=tmp_path,
                                    office_preview=False)
    assert result["status"] == "ready"
    source = result["sources"][0]
    assert source["metadata"]["physical_pages_known"] is False
    assert source["blocks"][0]["anchor"] == "word/body/1"
    assert source["blocks"][0]["part"] == "word/document.xml"
    assert any(block.get("repeated_header") for block in source["blocks"])
    assert any(block.get("page_break_before") for block in source["blocks"])
    nested = next(block for block in source["blocks"] if "Nested surcharge" in block["text"] and not block.get("structure_only"))
    assert nested["table"] and nested["row"] == 1 and nested["cell"] == 1
    assert nested["cell_anchor"] in nested["anchor"]
    assert sum("Merged service" in block["text"] for block in source["blocks"] if not block.get("structure_only")) == 1


def test_conversion_failure_preserves_uncomputed_formula_and_original(tmp_path):
    original = FIXTURES / "fees-multisheet.xlsx"
    digest = hashlib.sha256(original.read_bytes()).hexdigest()

    def fail(command, **kwargs):
        if "--version" in command:
            return SimpleNamespace(returncode=0, stdout="Synthetic converter contract", stderr="")
        return SimpleNamespace(returncode=1, stdout="", stderr="conversion failed")

    with patch.object(module.subprocess, "run", side_effect=fail):
        result = module.intake_document(original, output_dir=tmp_path,
                                        doc_converter="fixture-soffice", office_preview=False)
    assert result["status"] == "partial"
    assert hashlib.sha256(original.read_bytes()).hexdigest() == digest
    assert any(issue["code"] == "formula_recalculation_failed" for issue in result["issues"])
    assert any(issue["code"] == "formula_cache_missing" for issue in result["issues"])
    log = next(artifact for artifact in result["artifacts"] if artifact["kind"] == "conversion_log")
    assert json.loads(Path(log["path"]).read_text(encoding="utf-8"))["returncode"] == 1


def test_word_cell_without_optional_properties_retains_its_content(tmp_path):
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Handling Fee USD 50"
    cell._tc.remove(cell._tc.tcPr)
    path = tmp_path / "minimal-cell.docx"
    document.save(path)
    result = module.intake_document(path, output_dir=tmp_path / "intake", office_preview=False)
    assert result["status"] == "ready", result["issues"]
    assert "Handling Fee USD 50" in result["sources"][0]["text"]


@pytest.mark.skipif(not REAL_OFFICE, reason="requires configured project LibreOffice runtime")
def test_real_libreoffice_recalculates_xlsx_without_overwriting_original(tmp_path):
    original = FIXTURES / "fees-multisheet.xlsx"
    before = original.read_bytes()
    result = module.intake_document(original, output_dir=tmp_path, doc_converter=CONVERTER,
                                    office_preview=False)
    assert result["status"] == "ready", result["issues"]
    assert original.read_bytes() == before
    cells = {cell["coordinate"]: cell for block in result["sources"][0]["blocks"]
             if block.get("sheet") == "Charges" for cell in block["cells"]}
    assert [cells[cell]["value"] for cell in ["D5", "D6", "D7"]] == [50, 20, 70]
    assert cells["D5"]["formula"] == "=ROUND(B5*C5,2)"
    assert cells["D5"]["original_cached_value"] is None
    assert cells["D5"]["value_origin"] == "libreoffice-recalculation"
    assert any(artifact["kind"] == "office_formula-recalculation" for artifact in result["artifacts"])


@pytest.mark.skipif(not REAL_OFFICE, reason="requires configured project LibreOffice runtime")
def test_real_legacy_binary_xls_and_doc_preserve_sources_and_fee_text(tmp_path):
    for name in ["fees-multisheet.xls", "fees-cross-page.doc"]:
        original = FIXTURES / name
        assert original.read_bytes().startswith(b"\xd0\xcf\x11\xe0")
        before = original.read_bytes()
        result = module.intake_document(original, output_dir=tmp_path / original.suffix[1:],
                                        doc_converter=CONVERTER)
        assert result["status"] == "ready", result["issues"]
        assert original.read_bytes() == before
        assert any("Handling Fee" in source["text"] for source in result["sources"])
        assert all(artifact.get("original_content_id") for artifact in result["artifacts"]
                   if artifact["kind"].startswith("office_"))
        if original.suffix == ".doc":
            assert len(result["sources"]) == 2
            assert result["sources"][1]["parent_source_id"] == result["sources"][0]["source_id"]
            assert result["sources"][1]["metadata"]["office_preview"]["page_count"] > 1
        else:
            cells = {cell["coordinate"]: cell for block in result["sources"][0]["blocks"]
                     if block.get("sheet") == "Charges" for cell in block["cells"]}
            assert [cells[cell]["value"] for cell in ["D5", "D6", "D7"]] == [50, 20, 70]
            assert cells["D5"]["formula_origin"] == "libreoffice-conversion"


@pytest.mark.skipif(not REAL_OFFICE, reason="requires configured project LibreOffice runtime")
def test_real_docx_render_binds_preview_pages_to_its_hash(tmp_path):
    result = module.intake_document(FIXTURES / "fees-cross-page.docx", output_dir=tmp_path,
                                    doc_converter=CONVERTER)
    assert result["status"] == "ready", result["issues"]
    source = result["sources"][0]
    preview = source["metadata"]["office_preview"]
    assert preview["page_count"] > 1
    mapped = [block for block in source["blocks"] if "preview" in block]
    assert mapped and all(block["preview"]["artifact_sha256"] == preview["sha256"] for block in mapped)
    assert all("page" not in block for block in source["blocks"])
    assert source["metadata"]["physical_pages_known"] is False
