"""Create labelled synthetic Office files using format libraries and LibreOffice."""

from datetime import datetime
import argparse
import json
from pathlib import Path
import subprocess
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


def build(directory: Path, converter: str | None = None):
    directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Charges"
    sheet.merge_cells("A1:E1")
    sheet["A1"] = "INVOICE DEMO-OFFICE-001"
    sheet["A2"] = datetime(2026, 10, 21)
    sheet["A2"].number_format = "yyyy-mm-dd"
    for row in [
        ["Fee", "Quantity", "Unit Price", "Amount", "Currency"],
        ["Handling Fee", 2, 25, "=ROUND(B5*C5,2)", "USD"],
        ["Documentation Fee", 2, 10, "=B6*C6", "USD"],
        ["Total", None, None, "=SUM(D5:D6)", "USD"],
    ]:
        sheet.append(row)
    # append starts after row 2: insert a spacer so formulas retain their printed row IDs.
    sheet.insert_rows(3)
    sheet["D5"] = "=ROUND(B5*C5,2)"
    sheet["D6"] = "=B6*C6"
    sheet["D7"] = "=SUM(D5:D6)"
    sheet.row_dimensions[6].hidden = True
    sheet.column_dimensions["C"].hidden = True
    table = Table(displayName="FeeLines", ref="A4:E6")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    hidden = workbook.create_sheet("Additional charges")
    hidden.sheet_state = "hidden"
    hidden.append(["Fee", "Amount", "Currency"])
    hidden.append(["Storage Fee", 15, "EUR"])
    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_state = "veryHidden"
    metadata.append(["Invoice", "DEMO-OFFICE-001"])
    workbook.save(directory / "fees-multisheet.xlsx")
    workbook.close()

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(6), Inches(4)
    section.top_margin = section.bottom_margin = Inches(0.35)
    document.add_paragraph("INVOICE DEMO-OFFICE-001")
    document.add_paragraph("Issuer: Demo Logistics Ltd; Recipient: Demo Trading Ltd")
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Fee", "Amount", "Currency"]):
        cell.text = text
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for n in range(1, 26):
        for cell, text in zip(table.add_row().cells,
                              [f"Handling Fee line {n:02}", str(n), "USD"]):
            cell.text = text
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Merged service description USD 1"
    nested = table.cell(2, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested surcharge USD 3"
    last = document.add_paragraph("Explicit second section after page break")
    last.paragraph_format.page_break_before = True
    document.add_paragraph("Final amount USD 325")
    document.save(directory / "fees-cross-page.docx")

    manifest = dict(synthetic=True, description="Generated format fixtures, not exam answers",
                    invoice="DEMO-OFFICE-001", xlsx_amounts=[50, 20], xlsx_total=70,
                    hidden_amount=15, hidden_currency="EUR", word_rows=25)
    if converter:
        version = subprocess.run([converter, "--version"], check=True, capture_output=True,
                                 text=True, timeout=10).stdout.strip()
        manifest["converter_version"] = version
        for name, format_ in [("fees-multisheet.xlsx", "xls:MS Excel 97"),
                              ("fees-cross-page.docx", "doc:MS Word 97")]:
            with tempfile.TemporaryDirectory(prefix="office-fixture-") as profile:
                subprocess.run([converter, f"-env:UserInstallation={Path(profile).as_uri()}",
                                "--headless", "--convert-to", format_, "--outdir", str(directory),
                                str(directory / name)], check=True, capture_output=True,
                               text=True, timeout=120)
    (directory / "expected.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2),
                                              encoding="utf-8")


if __name__ == "__main__":
    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("--directory", type=Path, default=Path(__file__).parent)
    cli.add_argument("--converter")
    args = cli.parse_args()
    build(args.directory.resolve(), args.converter)
